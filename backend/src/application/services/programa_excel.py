"""Canonical Excel preview and import service for TASK-08.5."""

from __future__ import annotations

import io
import re
import uuid
from dataclasses import dataclass
from datetime import UTC, datetime
from typing import Protocol

from openpyxl import load_workbook

from src.application.dto.programa_documentos import StoredDocumentDTO
from src.application.dto.programa_excel import (
    ExcelCompetenciaPreviewDTO,
    ExcelConocimientoPreviewDTO,
    ExcelCriterioPreviewDTO,
    ExcelPendingAssignmentDTO,
    ExcelPendingSummaryDTO,
    ExcelPreviewSummaryDTO,
    ExcelProgramPreviewDTO,
    ExcelResultadoPreviewDTO,
    ExcelValidationIssueDTO,
    ProgramaExcelImportDTO,
    ProgramaExcelPreviewDTO,
)
from src.domain.drafts.types import TipoBloqueBorrador
from src.domain.shared.enums import (
    EstadoBloque,
    MotivoPendienteAsignacion,
    TipoConocimiento,
    TipoElementoCurricularPendiente,
)
from src.infrastructure.db.models.curriculum import (
    Competencia,
    Conocimiento,
    CriterioEvaluacion,
    ElementoCurricularPendiente,
    ProgramaFormacion,
    ResultadoAprendizaje,
)
from src.infrastructure.db.models.drafts import BorradorSesion

CANONICAL_SHEETS: dict[str, list[str]] = {
    "Programa": [
        "codigo_programa",
        "nombre_programa",
        "version_programa",
        "vigencia",
        "tipo_programa",
        "titulo_obtendra",
        "duracion_lectiva_horas",
        "duracion_productiva_horas",
        "duracion_total_horas",
        "fuente_archivo",
        "observaciones",
    ],
    "Competencias": [
        "competencia_id",
        "codigo_competencia",
        "nombre_competencia",
        "duracion_horas",
        "orden",
        "pagina_origen",
        "observaciones",
    ],
    "Resultados": [
        "competencia_id",
        "rap_id",
        "rap_numero",
        "resultado_aprendizaje",
        "orden",
        "pagina_origen",
        "observaciones",
    ],
    "Conocimientos": [
        "competencia_id",
        "rap_id",
        "tipo_conocimiento",
        "elemento_numero",
        "descripcion",
        "orden",
        "pagina_origen",
        "observaciones",
    ],
    "Criterios": [
        "competencia_id",
        "rap_id",
        "criterio_numero",
        "descripcion",
        "orden",
        "pagina_origen",
        "observaciones",
    ],
}

EXCEL_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ProgramaExcelDraftMissingError(Exception):
    """Raised when an Excel import has no current wizard draft."""


class InvalidProgramaExcelUploadError(Exception):
    """Raised when the upload is not an accepted canonical workbook."""


class ProgramaExcelValidationError(Exception):
    """Raised when a workbook cannot be imported after validation."""


class ProgramaExcelMissingPreviewError(Exception):
    """Raised when import confirmation has no valid preview to confirm."""


class ProgramaExcelStorageMissingError(Exception):
    """Raised when the stored workbook cannot be read for confirmation."""


@dataclass(frozen=True)
class ProgramaRow:
    codigo_programa: str
    nombre_programa: str
    version_programa: str | None
    raw: dict[str, object]


@dataclass(frozen=True)
class CompetenciaRow:
    competencia_id: str
    codigo_competencia: str
    nombre_competencia: str
    duracion_horas: int | None
    orden: int | None
    raw: dict[str, object]


@dataclass(frozen=True)
class ResultadoRow:
    competencia_id: str
    rap_id: str
    rap_numero: str | None
    resultado_aprendizaje: str
    orden: int | None
    raw: dict[str, object]


@dataclass(frozen=True)
class ConocimientoRow:
    competencia_id: str | None
    rap_id: str | None
    tipo_conocimiento: TipoConocimiento
    descripcion: str
    orden: int | None
    raw: dict[str, object]


@dataclass(frozen=True)
class CriterioRow:
    competencia_id: str | None
    rap_id: str | None
    descripcion: str
    orden: int | None
    raw: dict[str, object]


@dataclass(frozen=True)
class CanonicalWorkbook:
    programa: ProgramaRow | None
    competencias: list[CompetenciaRow]
    resultados: list[ResultadoRow]
    conocimientos: list[ConocimientoRow]
    criterios: list[CriterioRow]
    pendientes: list[ExcelPendingAssignmentDTO]
    errores: list[ExcelValidationIssueDTO]

    @property
    def is_valid(self) -> bool:
        """Return whether the workbook can proceed to preview or import."""
        return len(self.errores) == 0 and self.programa is not None


class DocumentStorageProtocol(Protocol):
    """Storage behavior required by the canonical Excel service."""

    async def save_excel(
        self,
        *,
        key: str,
        content: bytes,
        content_type: str,
        original_filename: str,
    ) -> StoredDocumentDTO:
        """Persist an Excel file and return object metadata."""

    async def read_excel(self, *, key: str) -> bytes:
        """Read a stored Excel file."""


class DraftRepositoryProtocol(Protocol):
    """Draft repository dependency used to associate imports to the wizard."""

    async def get_by_block_reference(
        self,
        tipo_bloque: TipoBloqueBorrador,
        referencia_id: uuid.UUID,
    ) -> BorradorSesion | None:
        """Return the current logical draft for a reference."""

    async def save(self, draft: BorradorSesion) -> BorradorSesion:
        """Persist modifications on an existing draft row."""


class AuditRepositoryProtocol(Protocol):
    """Audit dependency for relevant import events."""

    async def add_event(
        self,
        entidad: str,
        entidad_id: uuid.UUID,
        accion: str,
        detalle: dict[str, object] | None = None,
    ) -> object:
        """Persist a basic audit event."""


class ProgramaExcelRepositoryProtocol(Protocol):
    """Repository behavior required by the Excel import service."""

    async def get_programa(self, programa_id: uuid.UUID) -> ProgramaFormacion | None:
        """Return a program by id."""

    async def get_programa_by_code_version(
        self,
        codigo_programa: str,
        version_programa: str | None,
    ) -> ProgramaFormacion | None:
        """Return an existing logical program."""

    async def add_programa(
        self,
        *,
        codigo_programa: str,
        nombre_programa: str,
        version_programa: str | None,
    ) -> ProgramaFormacion:
        """Create a program row."""

    async def list_competencias(self, programa_id: uuid.UUID) -> list[Competencia]:
        """List program competences."""

    async def competencia_code_exists(
        self,
        *,
        programa_id: uuid.UUID,
        codigo_competencia: str,
    ) -> bool:
        """Return whether a competence code exists."""

    async def resultado_exists(
        self,
        *,
        competencia_id: uuid.UUID,
        descripcion: str,
    ) -> bool:
        """Return whether a result exists."""

    async def conocimiento_exists(
        self,
        *,
        competencia_id: uuid.UUID,
        tipo: TipoConocimiento,
        descripcion: str,
        resultado_id: uuid.UUID | None = None,
    ) -> bool:
        """Return whether a knowledge item exists."""

    async def criterio_exists(
        self,
        *,
        competencia_id: uuid.UUID,
        descripcion: str,
        resultado_id: uuid.UUID | None = None,
    ) -> bool:
        """Return whether a criterion exists."""

    async def add_competencia(
        self,
        *,
        programa_id: uuid.UUID,
        codigo_competencia: str,
        nombre_competencia: str,
        orden: int | None,
    ) -> Competencia:
        """Create a competence."""

    async def add_resultado(
        self,
        *,
        competencia_id: uuid.UUID,
        codigo_resultado: str | None,
        descripcion: str,
        orden: int | None,
    ) -> ResultadoAprendizaje:
        """Create a learning result."""

    async def add_conocimiento(
        self,
        *,
        competencia_id: uuid.UUID,
        tipo: TipoConocimiento,
        descripcion: str,
        orden: int | None,
        resultado_id: uuid.UUID | None = None,
    ) -> Conocimiento:
        """Create a knowledge item."""

    async def add_criterio(
        self,
        *,
        competencia_id: uuid.UUID,
        descripcion: str,
        orden: int | None,
        resultado_id: uuid.UUID | None = None,
    ) -> CriterioEvaluacion:
        """Create an evaluation criterion."""

    async def clear_pendientes(
        self,
        *,
        referencia_id: uuid.UUID,
    ) -> None:
        """Remove previous pending rows for a reference before reimporting."""

    async def add_pendiente(
        self,
        *,
        referencia_id: uuid.UUID,
        programa_id: uuid.UUID,
        pendiente: ExcelPendingAssignmentDTO,
        orden: int | None,
        raw_excel: dict[str, object] | None,
    ) -> ElementoCurricularPendiente:
        """Create a pending assignment row."""

    async def clear_curriculum(self, *, programa_id: uuid.UUID) -> None:
        """Remove all curriculum rows for a program before re-import."""


class AsyncSessionProtocol(Protocol):
    """Subset of async session behavior required by this service."""

    async def commit(self) -> None:
        """Commit the current transaction."""

    async def refresh(self, instance: object) -> None:
        """Refresh the provided ORM instance."""


class ProgramaExcelImportService:
    """Validate, preview and confirm canonical Excel curriculum imports."""

    def __init__(
        self,
        session: AsyncSessionProtocol,
        draft_repository: DraftRepositoryProtocol,
        audit_repository: AuditRepositoryProtocol,
        curriculum_repository: ProgramaExcelRepositoryProtocol,
        storage_service: DocumentStorageProtocol,
    ) -> None:
        """Initialize the service with explicit infrastructure ports."""
        self._session = session
        self._draft_repository = draft_repository
        self._audit_repository = audit_repository
        self._curriculum_repository = curriculum_repository
        self._storage_service = storage_service

    async def preview_program_excel(
        self,
        *,
        referencia_id: uuid.UUID,
        filename: str,
        content_type: str,
        content: bytes,
    ) -> ProgramaExcelPreviewDTO:
        """Store and validate a canonical Excel workbook without relational writes."""
        _validate_excel_upload(filename=filename, content=content)
        draft = await self._get_program_draft(referencia_id)
        workbook = parse_canonical_workbook(content)
        stored_document: StoredDocumentDTO | None = None

        if workbook.is_valid:
            stored_document = await self._storage_service.save_excel(
                key=_build_excel_storage_key(referencia_id, filename),
                content=content,
                content_type=content_type or EXCEL_CONTENT_TYPE,
                original_filename=filename,
            )

        preview = _build_preview_dto(
            referencia_id=referencia_id,
            workbook=workbook,
            document=stored_document,
        )
        draft.payload_json = _merge_excel_preview_into_payload(
            payload=draft.payload_json,
            preview=preview,
        )
        draft.paso_actual = "origen-documental"
        draft.estado_borrador = EstadoBloque.BORRADOR
        await self._draft_repository.save(draft)
        await self._audit_repository.add_event(
            entidad="BorradorSesion",
            entidad_id=draft.id,
            accion="EXCEL_PROGRAMA_PREVISUALIZADO",
            detalle={
                "referencia_id": str(referencia_id),
                "estado_validacion": preview.estado_validacion,
                "errores": len(preview.errores),
                "pendientes": preview.pendientes_resumen.total,
            },
        )
        await self._session.commit()
        await self._session.refresh(draft)
        return preview

    async def confirm_program_excel_import(
        self,
        *,
        referencia_id: uuid.UUID,
    ) -> ProgramaExcelImportDTO:
        """Materialize a previously validated canonical workbook."""
        draft = await self._get_program_draft(referencia_id)
        preview_payload = _get_valid_excel_preview_payload(draft.payload_json)
        document_payload = _as_record(preview_payload.get("documento"))
        storage_key = _read_string(document_payload or {}, "storage_key")
        if not storage_key:
            raise ProgramaExcelMissingPreviewError(
                "El borrador no tiene un Excel canonico validado para confirmar",
            )

        try:
            content = await self._storage_service.read_excel(key=storage_key)
        except FileNotFoundError as error:
            raise ProgramaExcelStorageMissingError(
                "El Excel canonico validado ya no existe en MinIO",
            ) from error

        workbook = parse_canonical_workbook(content)
        if not workbook.is_valid or workbook.programa is None:
            raise ProgramaExcelValidationError(
                "El Excel canonico almacenado ya no supera la validacion",
            )

        import_result = await self._materialize_workbook(
            referencia_id=referencia_id,
            draft=draft,
            workbook=workbook,
        )
        draft.payload_json = _merge_excel_import_into_payload(
            payload=draft.payload_json,
            result=import_result,
            competencias=await self._curriculum_repository.list_competencias(
                import_result.programa_id,
            ),
        )
        await self._draft_repository.save(draft)
        await self._audit_repository.add_event(
            entidad="ProgramaFormacion",
            entidad_id=import_result.programa_id,
            accion="EXCEL_CANONICO_IMPORTADO",
            detalle={
                "referencia_id": str(referencia_id),
                "competencias": import_result.resumen.competencias,
                "resultados": import_result.resumen.resultados,
                "conocimientos": import_result.resumen.conocimientos,
                "criterios": import_result.resumen.criterios,
                "pendientes": import_result.pendientes_resumen.total,
            },
        )
        await self._session.commit()
        await self._session.refresh(draft)
        return import_result

    async def _get_program_draft(self, referencia_id: uuid.UUID) -> BorradorSesion:
        draft = await self._draft_repository.get_by_block_reference(
            TipoBloqueBorrador.PROGRAMA,
            referencia_id,
        )
        if draft is None:
            raise ProgramaExcelDraftMissingError(
                "No existe un borrador de programa para asociar el Excel",
            )
        return draft

    async def _materialize_workbook(
        self,
        *,
        referencia_id: uuid.UUID,
        draft: BorradorSesion,
        workbook: CanonicalWorkbook,
    ) -> ProgramaExcelImportDTO:
        if workbook.programa is None:
            raise ProgramaExcelValidationError("La hoja Programa esta vacia")

        programa = await self._resolve_program(draft, workbook.programa)
        await self._curriculum_repository.clear_curriculum(programa_id=programa.id)
        competencia_by_excel_id: dict[str, Competencia] = {}
        resultado_by_excel_key: dict[tuple[str, str], uuid.UUID] = {}
        competencia_ids: list[uuid.UUID] = []
        resultado_ids: list[uuid.UUID] = []
        conocimiento_ids: list[uuid.UUID] = []
        criterio_ids: list[uuid.UUID] = []
        pendiente_ids: list[uuid.UUID] = []

        for row in workbook.competencias:
            if await self._curriculum_repository.competencia_code_exists(
                programa_id=programa.id,
                codigo_competencia=row.codigo_competencia,
            ):
                raise ProgramaExcelValidationError(
                    "Ya existe una competencia con codigo "
                    f"{row.codigo_competencia} en el programa actual",
                )
            competencia = await self._curriculum_repository.add_competencia(
                programa_id=programa.id,
                codigo_competencia=row.codigo_competencia,
                nombre_competencia=row.nombre_competencia,
                orden=row.orden,
            )
            await self._session.refresh(competencia)
            competencia_by_excel_id[row.competencia_id] = competencia
            competencia_ids.append(competencia.id)

        for resultado_row in workbook.resultados:
            competencia = competencia_by_excel_id[resultado_row.competencia_id]
            if await self._curriculum_repository.resultado_exists(
                competencia_id=competencia.id,
                descripcion=resultado_row.resultado_aprendizaje,
            ):
                raise ProgramaExcelValidationError(
                    "Ya existe un resultado duplicado en la competencia "
                    f"{resultado_row.competencia_id}",
                )
            resultado = await self._curriculum_repository.add_resultado(
                competencia_id=competencia.id,
                codigo_resultado=resultado_row.rap_id,
                descripcion=resultado_row.resultado_aprendizaje,
                orden=resultado_row.orden,
            )
            await self._session.refresh(resultado)
            resultado_ids.append(resultado.id)
            resultado_by_excel_key[
                (resultado_row.competencia_id, resultado_row.rap_id)
            ] = resultado.id

        for conocimiento_row in workbook.conocimientos:
            if conocimiento_row.competencia_id is None:
                raise ProgramaExcelValidationError(
                    "Conocimiento importable sin competencia resuelta",
                )
            competencia = competencia_by_excel_id[conocimiento_row.competencia_id]
            resultado_id = None
            if conocimiento_row.rap_id:
                resultado_id = resultado_by_excel_key.get(
                    (conocimiento_row.competencia_id, conocimiento_row.rap_id)
                )

            if await self._curriculum_repository.conocimiento_exists(
                competencia_id=competencia.id,
                tipo=conocimiento_row.tipo_conocimiento,
                descripcion=conocimiento_row.descripcion,
                resultado_id=resultado_id,
            ):
                raise ProgramaExcelValidationError(
                    "Ya existe un conocimiento duplicado en la competencia "
                    f"{conocimiento_row.competencia_id}",
                )
            conocimiento = await self._curriculum_repository.add_conocimiento(
                competencia_id=competencia.id,
                tipo=conocimiento_row.tipo_conocimiento,
                descripcion=conocimiento_row.descripcion,
                orden=conocimiento_row.orden,
                resultado_id=resultado_id,
            )
            await self._session.refresh(conocimiento)
            conocimiento_ids.append(conocimiento.id)

        for criterio_row in workbook.criterios:
            if criterio_row.competencia_id is None:
                raise ProgramaExcelValidationError(
                    "Criterio importable sin competencia resuelta",
                )
            competencia = competencia_by_excel_id[criterio_row.competencia_id]
            resultado_id = None
            if criterio_row.rap_id:
                resultado_id = resultado_by_excel_key.get(
                    (criterio_row.competencia_id, criterio_row.rap_id)
                )

            if await self._curriculum_repository.criterio_exists(
                competencia_id=competencia.id,
                descripcion=criterio_row.descripcion,
                resultado_id=resultado_id,
            ):
                raise ProgramaExcelValidationError(
                    "Ya existe un criterio duplicado en la competencia "
                    f"{criterio_row.competencia_id}",
                )
            criterio = await self._curriculum_repository.add_criterio(
                competencia_id=competencia.id,
                descripcion=criterio_row.descripcion,
                orden=criterio_row.orden,
                resultado_id=resultado_id,
            )
            await self._session.refresh(criterio)
            criterio_ids.append(criterio.id)

        await self._curriculum_repository.clear_pendientes(
            referencia_id=referencia_id,
        )
        for index, pendiente in enumerate(workbook.pendientes, start=1):
            raw_excel = _raw_excel_for_pending(workbook, pendiente)
            pending_row = await self._curriculum_repository.add_pendiente(
                referencia_id=referencia_id,
                programa_id=programa.id,
                pendiente=pendiente,
                orden=index,
                raw_excel=raw_excel,
            )
            await self._session.refresh(pending_row)
            pendiente_ids.append(pending_row.id)

        return ProgramaExcelImportDTO(
            referencia_id=referencia_id,
            programa_id=programa.id,
            competencia_ids=competencia_ids,
            resultado_ids=resultado_ids,
            conocimiento_ids=conocimiento_ids,
            criterio_ids=criterio_ids,
            pendiente_ids=pendiente_ids,
            resumen=_build_summary(workbook),
            pendientes_resumen=_build_pending_summary(workbook.pendientes),
        )

    async def _resolve_program(
        self,
        draft: BorradorSesion,
        row: ProgramaRow,
    ) -> ProgramaFormacion:
        draft_programa_id = _extract_programa_id(draft.payload_json)
        if draft_programa_id is not None:
            programa = await self._curriculum_repository.get_programa(
                draft_programa_id,
            )
            if programa is None:
                raise ProgramaExcelValidationError(
                    "El borrador referencia un programa que ya no existe",
                )
            if (
                programa.codigo_programa.lower() != row.codigo_programa.lower()
                or programa.version_programa != row.version_programa
            ):
                raise ProgramaExcelValidationError(
                    "El Excel no coincide con el programa ya asociado al borrador",
                )
            return programa

        programa = await self._curriculum_repository.get_programa_by_code_version(
            row.codigo_programa,
            row.version_programa,
        )
        if programa is not None:
            return programa

        programa = await self._curriculum_repository.add_programa(
            codigo_programa=row.codigo_programa,
            nombre_programa=row.nombre_programa,
            version_programa=row.version_programa,
        )
        await self._session.refresh(programa)
        return programa


def parse_canonical_workbook(content: bytes) -> CanonicalWorkbook:
    """Read and validate the canonical workbook contract."""
    errores: list[ExcelValidationIssueDTO] = []
    try:
        workbook = load_workbook(
            filename=io.BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except Exception as error:
        return CanonicalWorkbook(
            programa=None,
            competencias=[],
            resultados=[],
            conocimientos=[],
            criterios=[],
            pendientes=[],
            errores=[
                ExcelValidationIssueDTO(
                    hoja="Workbook",
                    fila=None,
                    campo=None,
                    mensaje=f"No fue posible leer el Excel canonico: {error}",
                ),
            ],
        )

    missing_sheets = [
        sheet_name
        for sheet_name in CANONICAL_SHEETS
        if sheet_name not in workbook.sheetnames
    ]
    for sheet_name in missing_sheets:
        errores.append(
            ExcelValidationIssueDTO(
                hoja=sheet_name,
                fila=None,
                campo=None,
                mensaje="Hoja obligatoria faltante",
            ),
        )
    if missing_sheets:
        return CanonicalWorkbook(None, [], [], [], [], [], errores)

    rows_by_sheet: dict[str, list[dict[str, object]]] = {}
    for sheet_name, expected_headers in CANONICAL_SHEETS.items():
        sheet = workbook[sheet_name]
        header_values = [
            _cell_to_string(value)
            for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        ]
        if header_values != expected_headers:
            errores.append(
                ExcelValidationIssueDTO(
                    hoja=sheet_name,
                    fila=1,
                    campo=None,
                    mensaje="Encabezados invalidos o en orden distinto al canonico",
                ),
            )
            continue

        sheet_rows: list[dict[str, object]] = []
        for row_index, values in enumerate(
            sheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            if all(_cell_to_string(value) == "" for value in values):
                continue
            record: dict[str, object] = {
                header: value
                for header, value in zip(expected_headers, values, strict=False)
            }
            record["_row_index"] = row_index
            sheet_rows.append(record)
        rows_by_sheet[sheet_name] = sheet_rows

    if errores:
        return CanonicalWorkbook(None, [], [], [], [], [], errores)

    programa = _parse_programa(rows_by_sheet["Programa"], errores)
    competencias = _parse_competencias(rows_by_sheet["Competencias"], errores)
    resultados = _parse_resultados(rows_by_sheet["Resultados"], errores)
    conocimientos = _parse_conocimientos(rows_by_sheet["Conocimientos"], errores)
    criterios = _parse_criterios(rows_by_sheet["Criterios"], errores)
    conocimientos, criterios, pendientes = _split_assignable_rows(
        competencias=competencias,
        resultados=resultados,
        conocimientos=conocimientos,
        criterios=criterios,
    )

    _validate_cross_references(
        competencias=competencias,
        resultados=resultados,
        conocimientos=conocimientos,
        criterios=criterios,
        errores=errores,
    )
    _validate_duplicates(
        competencias=competencias,
        resultados=resultados,
        conocimientos=conocimientos,
        criterios=criterios,
        pendientes=pendientes,
        errores=errores,
    )

    return CanonicalWorkbook(
        programa=programa,
        competencias=competencias,
        resultados=resultados,
        conocimientos=conocimientos,
        criterios=criterios,
        pendientes=pendientes,
        errores=errores,
    )


def _parse_programa(
    rows: list[dict[str, object]],
    errores: list[ExcelValidationIssueDTO],
) -> ProgramaRow | None:
    if len(rows) != 1:
        errores.append(
            ExcelValidationIssueDTO(
                hoja="Programa",
                fila=None,
                campo=None,
                mensaje="La hoja Programa debe tener exactamente una fila de datos",
            ),
        )
        return None

    row = rows[0]
    codigo = _required_string(row, "Programa", "codigo_programa", errores)
    nombre = _required_string(row, "Programa", "nombre_programa", errores)
    version = _optional_string(row.get("version_programa"))
    _optional_int(row, "Programa", "duracion_lectiva_horas", errores)
    _optional_int(row, "Programa", "duracion_productiva_horas", errores)
    _optional_int(row, "Programa", "duracion_total_horas", errores)
    if codigo is None or nombre is None:
        return None
    return ProgramaRow(codigo, nombre, version, _public_record(row))


def _parse_competencias(
    rows: list[dict[str, object]],
    errores: list[ExcelValidationIssueDTO],
) -> list[CompetenciaRow]:
    parsed: list[CompetenciaRow] = []
    for row in rows:
        competencia_id = _required_string(
            row,
            "Competencias",
            "competencia_id",
            errores,
        )
        codigo = _required_string(
            row,
            "Competencias",
            "codigo_competencia",
            errores,
        )
        nombre = _required_string(
            row,
            "Competencias",
            "nombre_competencia",
            errores,
        )
        duracion = _optional_int(row, "Competencias", "duracion_horas", errores)
        orden = _optional_int(row, "Competencias", "orden", errores)
        if competencia_id is None or codigo is None or nombre is None:
            continue
        parsed.append(
            CompetenciaRow(
                competencia_id=competencia_id,
                codigo_competencia=codigo,
                nombre_competencia=nombre,
                duracion_horas=duracion,
                orden=orden,
                raw=_public_record(row),
            ),
        )
    return parsed


def _parse_resultados(
    rows: list[dict[str, object]],
    errores: list[ExcelValidationIssueDTO],
) -> list[ResultadoRow]:
    parsed: list[ResultadoRow] = []
    for row in rows:
        competencia_id = _required_string(
            row,
            "Resultados",
            "competencia_id",
            errores,
        )
        rap_id = _required_string(row, "Resultados", "rap_id", errores)
        descripcion = _required_string(
            row,
            "Resultados",
            "resultado_aprendizaje",
            errores,
        )
        orden = _optional_int(row, "Resultados", "orden", errores)
        if competencia_id is None or rap_id is None or descripcion is None:
            continue
        parsed.append(
            ResultadoRow(
                competencia_id=competencia_id,
                rap_id=rap_id,
                rap_numero=_optional_string(row.get("rap_numero")),
                resultado_aprendizaje=descripcion,
                orden=orden,
                raw=_public_record(row),
            ),
        )
    return parsed


def _parse_conocimientos(
    rows: list[dict[str, object]],
    errores: list[ExcelValidationIssueDTO],
) -> list[ConocimientoRow]:
    parsed: list[ConocimientoRow] = []
    for row in rows:
        competencia_id = _optional_string(row.get("competencia_id"))
        raw_tipo = _required_string(
            row,
            "Conocimientos",
            "tipo_conocimiento",
            errores,
        )
        descripcion = _required_string(
            row,
            "Conocimientos",
            "descripcion",
            errores,
        )
        orden = _optional_int(row, "Conocimientos", "orden", errores)
        tipo = _parse_tipo_conocimiento(raw_tipo, row, errores)
        if tipo is None or descripcion is None:
            continue
        parsed.append(
            ConocimientoRow(
                competencia_id=competencia_id,
                rap_id=_optional_string(row.get("rap_id")),
                tipo_conocimiento=tipo,
                descripcion=descripcion,
                orden=orden,
                raw=_public_record(row),
            ),
        )
    return parsed


def _parse_criterios(
    rows: list[dict[str, object]],
    errores: list[ExcelValidationIssueDTO],
) -> list[CriterioRow]:
    parsed: list[CriterioRow] = []
    for row in rows:
        competencia_id = _optional_string(row.get("competencia_id"))
        descripcion = _required_string(row, "Criterios", "descripcion", errores)
        orden = _optional_int(row, "Criterios", "orden", errores)
        if descripcion is None:
            continue
        parsed.append(
            CriterioRow(
                competencia_id=competencia_id,
                rap_id=_optional_string(row.get("rap_id")),
                descripcion=descripcion,
                orden=orden,
                raw=_public_record(row),
            ),
        )
    return parsed


def _split_assignable_rows(
    *,
    competencias: list[CompetenciaRow],
    resultados: list[ResultadoRow],
    conocimientos: list[ConocimientoRow],
    criterios: list[CriterioRow],
) -> tuple[list[ConocimientoRow], list[CriterioRow], list[ExcelPendingAssignmentDTO]]:
    competencia_ids = {row.competencia_id for row in competencias}
    pending: list[ExcelPendingAssignmentDTO] = []
    assignable_conocimientos: list[ConocimientoRow] = []
    assignable_criterios: list[CriterioRow] = []

    for conocimiento_row in conocimientos:
        motivo = _pending_reason_for_row(
            competencia_id=conocimiento_row.competencia_id,
            competencia_ids=competencia_ids,
        )
        if motivo is not None:
            pending.append(_pending_from_conocimiento(conocimiento_row, motivo))
            continue
        assignable_conocimientos.append(conocimiento_row)

    for criterio_row in criterios:
        motivo = _pending_reason_for_row(
            competencia_id=criterio_row.competencia_id,
            competencia_ids=competencia_ids,
        )
        if motivo is not None:
            pending.append(_pending_from_criterio(criterio_row, motivo))
            continue
        assignable_criterios.append(criterio_row)

    assignable_conocimientos, duplicated_conocimientos = (
        _move_duplicate_conocimientos_to_pending(assignable_conocimientos)
    )
    assignable_criterios, duplicated_criterios = _move_duplicate_criterios_to_pending(
        assignable_criterios
    )
    pending.extend(duplicated_conocimientos)
    pending.extend(duplicated_criterios)
    return assignable_conocimientos, assignable_criterios, pending


def _pending_reason_for_row(
    *,
    competencia_id: str | None,
    competencia_ids: set[str],
) -> MotivoPendienteAsignacion | None:
    if competencia_id is None or competencia_id not in competencia_ids:
        return MotivoPendienteAsignacion.COMPETENCIA_NO_IDENTIFICADA
    return None


def _move_duplicate_conocimientos_to_pending(
    rows: list[ConocimientoRow],
) -> tuple[list[ConocimientoRow], list[ExcelPendingAssignmentDTO]]:
    seen: set[tuple[str | None, str, str]] = set()
    assignable: list[ConocimientoRow] = []
    pending: list[ExcelPendingAssignmentDTO] = []
    for row in rows:
        key = (
            row.competencia_id,
            row.tipo_conocimiento.value,
            _norm(row.descripcion),
        )
        if key in seen:
            pending.append(
                _pending_from_conocimiento(
                    row,
                    MotivoPendienteAsignacion.ASOCIACION_AMBIGUA,
                )
            )
            continue
        seen.add(key)
        assignable.append(row)
    return assignable, pending


def _move_duplicate_criterios_to_pending(
    rows: list[CriterioRow],
) -> tuple[list[CriterioRow], list[ExcelPendingAssignmentDTO]]:
    seen: set[tuple[str | None, str]] = set()
    assignable: list[CriterioRow] = []
    pending: list[ExcelPendingAssignmentDTO] = []
    for row in rows:
        key = (row.competencia_id, _norm(row.descripcion))
        if key in seen:
            pending.append(
                _pending_from_criterio(
                    row,
                    MotivoPendienteAsignacion.ASOCIACION_AMBIGUA,
                )
            )
            continue
        seen.add(key)
        assignable.append(row)
    return assignable, pending


def _pending_from_conocimiento(
    row: ConocimientoRow,
    motivo: MotivoPendienteAsignacion,
) -> ExcelPendingAssignmentDTO:
    return ExcelPendingAssignmentDTO(
        tipo_elemento=TipoElementoCurricularPendiente.CONOCIMIENTO,
        tipo_conocimiento=row.tipo_conocimiento,
        descripcion=row.descripcion,
        competencia_id_origen_excel=row.competencia_id,
        rap_id_origen_excel=row.rap_id,
        motivo=motivo,
        hoja="Conocimientos",
        fila=_row_index(row.raw),
        raw_excel=row.raw,
    )


def _pending_from_criterio(
    row: CriterioRow,
    motivo: MotivoPendienteAsignacion,
) -> ExcelPendingAssignmentDTO:
    return ExcelPendingAssignmentDTO(
        tipo_elemento=TipoElementoCurricularPendiente.CRITERIO,
        tipo_conocimiento=None,
        descripcion=row.descripcion,
        competencia_id_origen_excel=row.competencia_id,
        rap_id_origen_excel=row.rap_id,
        motivo=motivo,
        hoja="Criterios",
        fila=_row_index(row.raw),
        raw_excel=row.raw,
    )


def _raw_excel_for_pending(
    workbook: CanonicalWorkbook,
    pendiente: ExcelPendingAssignmentDTO,
) -> dict[str, object] | None:
    del workbook
    return pendiente.raw_excel


def _validate_cross_references(
    *,
    competencias: list[CompetenciaRow],
    resultados: list[ResultadoRow],
    conocimientos: list[ConocimientoRow],
    criterios: list[CriterioRow],
    errores: list[ExcelValidationIssueDTO],
) -> None:
    competencia_ids = {row.competencia_id for row in competencias}

    for resultado_row in resultados:
        _validate_competencia_reference(resultado_row, competencia_ids, errores)
    for conocimiento_row in conocimientos:
        _validate_competencia_reference(conocimiento_row, competencia_ids, errores)
    for criterio_row in criterios:
        _validate_competencia_reference(criterio_row, competencia_ids, errores)


def _validate_competencia_reference(
    row: ResultadoRow | ConocimientoRow | CriterioRow,
    competencia_ids: set[str],
    errores: list[ExcelValidationIssueDTO],
) -> None:
    if row.competencia_id in competencia_ids:
        return
    errores.append(
        ExcelValidationIssueDTO(
            hoja=_sheet_for_row(row),
            fila=_row_index(row.raw),
            campo="competencia_id",
            mensaje="competencia_id no existe en hoja Competencias",
        ),
    )


def _validate_rap_reference(
    row: ConocimientoRow | CriterioRow,
    rap_keys: set[tuple[str, str]],
    errores: list[ExcelValidationIssueDTO],
) -> None:
    if row.rap_id is None or (row.competencia_id, row.rap_id) in rap_keys:
        return
    errores.append(
        ExcelValidationIssueDTO(
            hoja=_sheet_for_row(row),
            fila=_row_index(row.raw),
            campo="rap_id",
            mensaje="rap_id no existe en Resultados para la competencia",
        ),
    )


def _validate_duplicates(
    *,
    competencias: list[CompetenciaRow],
    resultados: list[ResultadoRow],
    conocimientos: list[ConocimientoRow],
    criterios: list[CriterioRow],
    pendientes: list[ExcelPendingAssignmentDTO],
    errores: list[ExcelValidationIssueDTO],
) -> None:
    del conocimientos, criterios, pendientes
    _reject_duplicate_keys(
        "Competencias",
        [(row.codigo_competencia.lower(), row) for row in competencias],
        "codigo_competencia",
        "Competencia duplicada por codigo dentro del programa",
        errores,
    )
    _reject_duplicate_keys(
        "Resultados",
        [((row.competencia_id, row.rap_id), row) for row in resultados],
        "rap_id",
        "Resultado duplicado por rap_id dentro de la competencia",
        errores,
        labels=("competencia_id", "rap_id"),
    )
    _reject_duplicate_keys(
        "Resultados",
        [
            ((row.competencia_id, _norm(row.resultado_aprendizaje)), row)
            for row in resultados
        ],
        "resultado_aprendizaje",
        "Resultado duplicado exacto dentro de la competencia",
        errores,
        labels=("competencia_id", "resultado_aprendizaje"),
    )


def _reject_duplicate_keys(
    hoja: str,
    keys: list[tuple[object, object]],
    campo: str,
    mensaje: str,
    errores: list[ExcelValidationIssueDTO],
    labels: tuple[str, ...] | None = None,
) -> None:
    seen: dict[object, int | None] = {}
    for key, row in keys:
        if key in seen:
            raw = getattr(row, "raw", {})
            errores.append(
                ExcelValidationIssueDTO(
                    hoja=hoja,
                    fila=_row_index(raw),
                    campo=campo,
                    mensaje=_duplicate_message(
                        mensaje=mensaje,
                        key=key,
                        labels=labels,
                        original_row=seen[key],
                    ),
                ),
            )
            continue
        raw = getattr(row, "raw", {})
        seen[key] = _row_index(raw)


def _duplicate_message(
    *,
    mensaje: str,
    key: object,
    labels: tuple[str, ...] | None,
    original_row: int | None,
) -> str:
    details: list[str] = []
    if labels is not None and isinstance(key, tuple):
        details.extend(
            f"{label}={value or '<vacio>'}"
            for label, value in zip(labels, key, strict=False)
        )
    if original_row is not None:
        details.append(f"fila_original={original_row}")
    if not details:
        return mensaje
    return f"{mensaje} ({', '.join(details)})"


def _sheet_for_row(row: object) -> str:
    if isinstance(row, ResultadoRow):
        return "Resultados"
    if isinstance(row, ConocimientoRow):
        return "Conocimientos"
    if isinstance(row, CriterioRow):
        return "Criterios"
    if isinstance(row, CompetenciaRow):
        return "Competencias"
    return "Workbook"


def _validate_excel_upload(*, filename: str, content: bytes) -> None:
    if not filename.strip():
        raise InvalidProgramaExcelUploadError("El archivo Excel es obligatorio")
    if not filename.lower().endswith(".xlsx"):
        raise InvalidProgramaExcelUploadError("Solo se aceptan archivos .xlsx")
    if not content:
        raise InvalidProgramaExcelUploadError("El Excel seleccionado esta vacio")


def _build_preview_dto(
    *,
    referencia_id: uuid.UUID,
    workbook: CanonicalWorkbook,
    document: StoredDocumentDTO | None,
) -> ProgramaExcelPreviewDTO:
    rap_keys = {(row.competencia_id, row.rap_id) for row in workbook.resultados}
    return ProgramaExcelPreviewDTO(
        referencia_id=referencia_id,
        documento=document,
        valid=workbook.is_valid,
        estado_validacion="VALIDO" if workbook.is_valid else "INVALIDO",
        resumen=_build_summary(workbook),
        programa=(
            ExcelProgramPreviewDTO(
                codigo_programa=workbook.programa.codigo_programa,
                nombre_programa=workbook.programa.nombre_programa,
                version_programa=workbook.programa.version_programa,
            )
            if workbook.programa is not None
            else None
        ),
        competencias=[
            ExcelCompetenciaPreviewDTO(
                competencia_id=row.competencia_id,
                codigo_competencia=row.codigo_competencia,
                nombre_competencia=row.nombre_competencia,
                resultados=sum(
                    1
                    for item in workbook.resultados
                    if item.competencia_id == row.competencia_id
                ),
                conocimientos=sum(
                    1
                    for item in workbook.conocimientos
                    if item.competencia_id == row.competencia_id
                ),
                criterios=sum(
                    1
                    for item in workbook.criterios
                    if item.competencia_id == row.competencia_id
                ),
                resultados_detalle=[
                    ExcelResultadoPreviewDTO(
                        rap_id=item.rap_id,
                        rap_numero=item.rap_numero,
                        descripcion=item.resultado_aprendizaje,
                    )
                    for item in workbook.resultados
                    if item.competencia_id == row.competencia_id
                ],
                conocimientos_detalle=[
                    ExcelConocimientoPreviewDTO(
                        tipo_conocimiento=item.tipo_conocimiento,
                        descripcion=item.descripcion,
                        rap_id=_resolved_rap_id(item, rap_keys),
                    )
                    for item in workbook.conocimientos
                    if item.competencia_id == row.competencia_id
                ],
                criterios_detalle=[
                    ExcelCriterioPreviewDTO(
                        descripcion=item.descripcion,
                        rap_id=_resolved_rap_id(item, rap_keys),
                    )
                    for item in workbook.criterios
                    if item.competencia_id == row.competencia_id
                ],
            )
            for row in workbook.competencias
        ],
        pendientes_resumen=_build_pending_summary(workbook.pendientes),
        pendientes=workbook.pendientes,
        errores=workbook.errores,
    )


def _resolved_rap_id(
    row: ConocimientoRow | CriterioRow,
    rap_keys: set[tuple[str, str]],
) -> str | None:
    if row.competencia_id is None or row.rap_id is None:
        return None
    return row.rap_id if (row.competencia_id, row.rap_id) in rap_keys else None


def _build_summary(workbook: CanonicalWorkbook) -> ExcelPreviewSummaryDTO:
    return ExcelPreviewSummaryDTO(
        programa=1 if workbook.programa is not None else 0,
        competencias=len(workbook.competencias),
        resultados=len(workbook.resultados),
        conocimientos=len(workbook.conocimientos),
        criterios=len(workbook.criterios),
    )


def _build_pending_summary(
    pendientes: list[ExcelPendingAssignmentDTO],
) -> ExcelPendingSummaryDTO:
    return ExcelPendingSummaryDTO(
        total=len(pendientes),
        conocimientos=sum(
            1
            for item in pendientes
            if item.tipo_elemento == TipoElementoCurricularPendiente.CONOCIMIENTO
        ),
        criterios=sum(
            1
            for item in pendientes
            if item.tipo_elemento == TipoElementoCurricularPendiente.CRITERIO
        ),
    )


def _merge_excel_preview_into_payload(
    *,
    payload: dict[str, object],
    preview: ProgramaExcelPreviewDTO,
) -> dict[str, object]:
    next_payload = dict(payload)
    now = datetime.now(UTC).isoformat()
    meta = _as_record(next_payload.get("meta")) or {}
    next_payload["meta"] = {
        **meta,
        "entryMode": "EXCEL",
        "lastInteractionAt": now,
        "touchedSteps": _append_touched_step(
            meta.get("touchedSteps"),
            "origen-documental",
        ),
    }
    documental = _as_record(next_payload.get("documental")) or {}
    documental["programa_excel"] = _preview_to_payload(preview, updated_at=now)
    next_payload["documental"] = documental
    return next_payload


def _merge_excel_import_into_payload(
    *,
    payload: dict[str, object],
    result: ProgramaExcelImportDTO,
    competencias: list[Competencia],
) -> dict[str, object]:
    next_payload = dict(payload)
    now = datetime.now(UTC).isoformat()
    documental = _as_record(next_payload.get("documental")) or {}
    excel_payload = _as_record(documental.get("programa_excel")) or {}
    excel_payload["confirmacion"] = {
        "estado": "IMPORTADO",
        "confirmed_at": now,
        "programa_id": str(result.programa_id),
        "competencia_ids": [str(item) for item in result.competencia_ids],
        "resultado_ids": [str(item) for item in result.resultado_ids],
        "conocimiento_ids": [str(item) for item in result.conocimiento_ids],
        "criterio_ids": [str(item) for item in result.criterio_ids],
        "pendiente_ids": [str(item) for item in result.pendiente_ids],
        "pendientes_resumen": _pending_summary_to_payload(
            result.pendientes_resumen,
        ),
    }
    documental["programa_excel"] = excel_payload
    next_payload["documental"] = documental

    if result.resumen.programa == 1:
        preview = _as_record(excel_payload.get("preview")) or {}
        programa = _as_record(preview.get("programa")) or {}
        next_payload["programa"] = {
            **(_as_record(next_payload.get("programa")) or {}),
            "codigo_programa": _read_string(programa, "codigo_programa"),
            "nombre_programa": _read_string(programa, "nombre_programa"),
            "version_programa": _read_string(programa, "version_programa"),
        }

    next_payload["curricular"] = {
        "programa_formacion_id": str(result.programa_id),
        "competencias": [_competencia_payload_item(item) for item in competencias],
        "excel_import": {
            "estado": "IMPORTADO",
            "updated_at": now,
            "resumen": _summary_to_payload(result.resumen),
            "pendientes_resumen": _pending_summary_to_payload(
                result.pendientes_resumen,
            ),
        },
    }
    return next_payload


def _preview_to_payload(
    preview: ProgramaExcelPreviewDTO,
    *,
    updated_at: str,
) -> dict[str, object]:
    return {
        "documento": _document_to_payload(preview.documento),
        "preview": {
            "valid": preview.valid,
            "estado_validacion": preview.estado_validacion,
            "resumen": _summary_to_payload(preview.resumen),
            "programa": (
                {
                    "codigo_programa": preview.programa.codigo_programa,
                    "nombre_programa": preview.programa.nombre_programa,
                    "version_programa": preview.programa.version_programa,
                }
                if preview.programa is not None
                else None
            ),
            "competencias": [
                {
                    "competencia_id": item.competencia_id,
                    "codigo_competencia": item.codigo_competencia,
                    "nombre_competencia": item.nombre_competencia,
                    "resultados": item.resultados,
                    "conocimientos": item.conocimientos,
                    "criterios": item.criterios,
                    "resultados_detalle": [
                        {
                            "rap_id": resultado.rap_id,
                            "rap_numero": resultado.rap_numero,
                            "descripcion": resultado.descripcion,
                        }
                        for resultado in item.resultados_detalle
                    ],
                    "conocimientos_detalle": [
                        {
                            "tipo_conocimiento": conocimiento.tipo_conocimiento.value,
                            "descripcion": conocimiento.descripcion,
                            "rap_id": conocimiento.rap_id,
                        }
                        for conocimiento in item.conocimientos_detalle
                    ],
                    "criterios_detalle": [
                        {
                            "descripcion": criterio.descripcion,
                            "rap_id": criterio.rap_id,
                        }
                        for criterio in item.criterios_detalle
                    ],
                }
                for item in preview.competencias
            ],
            "pendientes_resumen": _pending_summary_to_payload(
                preview.pendientes_resumen,
            ),
            "pendientes": [
                {
                    "tipo_elemento": item.tipo_elemento.value,
                    "tipo_conocimiento": (
                        item.tipo_conocimiento.value
                        if item.tipo_conocimiento is not None
                        else None
                    ),
                    "descripcion": item.descripcion,
                    "competencia_id_origen_excel": item.competencia_id_origen_excel,
                    "rap_id_origen_excel": item.rap_id_origen_excel,
                    "motivo": item.motivo.value,
                    "hoja": item.hoja,
                    "fila": item.fila,
                }
                for item in preview.pendientes
            ],
            "errores": [
                {
                    "hoja": issue.hoja,
                    "fila": issue.fila,
                    "campo": issue.campo,
                    "mensaje": issue.mensaje,
                }
                for issue in preview.errores
            ],
        },
        "confirmacion": {"estado": "PENDIENTE"},
        "updated_at": updated_at,
    }


def _document_to_payload(
    document: StoredDocumentDTO | None,
) -> dict[str, object] | None:
    if document is None:
        return None
    return {
        "original_filename": document.original_filename,
        "storage_key": document.storage_key,
        "size_bytes": document.size_bytes,
        "content_type": document.content_type,
        "checksum_sha256": document.checksum_sha256,
        "etag": document.etag,
    }


def _summary_to_payload(summary: ExcelPreviewSummaryDTO) -> dict[str, int]:
    return {
        "programa": summary.programa,
        "competencias": summary.competencias,
        "resultados": summary.resultados,
        "conocimientos": summary.conocimientos,
        "criterios": summary.criterios,
    }


def _pending_summary_to_payload(summary: ExcelPendingSummaryDTO) -> dict[str, int]:
    return {
        "total": summary.total,
        "conocimientos": summary.conocimientos,
        "criterios": summary.criterios,
    }


def _get_valid_excel_preview_payload(
    payload: dict[str, object],
) -> dict[str, object]:
    documental = _as_record(payload.get("documental")) or {}
    excel = _as_record(documental.get("programa_excel"))
    preview = _as_record(excel.get("preview") if excel is not None else None)
    if excel is None or preview is None or preview.get("valid") is not True:
        raise ProgramaExcelMissingPreviewError(
            "El borrador no tiene un preview Excel valido para confirmar",
        )
    confirmation = _as_record(excel.get("confirmacion")) or {}
    if confirmation.get("estado") == "IMPORTADO":
        raise ProgramaExcelValidationError(
            "El Excel canonico ya fue confirmado en este borrador",
        )
    return excel


def _extract_programa_id(payload: dict[str, object]) -> uuid.UUID | None:
    curricular = _as_record(payload.get("curricular"))
    raw_programa_id = curricular.get("programa_formacion_id") if curricular else None
    if not isinstance(raw_programa_id, str) or not raw_programa_id:
        return None
    try:
        return uuid.UUID(raw_programa_id)
    except ValueError:
        return None


def _competencia_payload_item(competencia: Competencia) -> dict[str, object]:
    return {
        "id": str(competencia.id),
        "programa_id": str(competencia.programa_id),
        "codigo_competencia": competencia.codigo_competencia,
        "nombre_competencia": competencia.nombre_competencia,
        "orden": competencia.orden,
        "estado": competencia.estado.value,
        "origen_campo": competencia.origen_campo.value,
        "fecha_creacion": competencia.fecha_creacion.isoformat(),
        "fecha_actualizacion": competencia.fecha_actualizacion.isoformat(),
        "resultados": [
            {
                "id": str(item.id),
                "competencia_id": str(item.competencia_id),
                "codigo_resultado": item.codigo_resultado,
                "descripcion": item.descripcion,
                "orden": item.orden,
                "estado": item.estado.value,
                "fecha_creacion": item.fecha_creacion.isoformat(),
                "fecha_actualizacion": item.fecha_actualizacion.isoformat(),
            }
            for item in sorted(
                competencia.resultados,
                key=lambda item: (item.orden is None, item.orden or 0),
            )
        ],
        "conocimientos": [
            {
                "id": str(item.id),
                "competencia_id": str(item.competencia_id),
                "resultado_id": str(item.resultado_id) if item.resultado_id else None,
                "tipo": item.tipo.value,
                "descripcion": item.descripcion,
                "orden": item.orden,
                "estado": item.estado.value,
                "fecha_creacion": item.fecha_creacion.isoformat(),
                "fecha_actualizacion": item.fecha_actualizacion.isoformat(),
            }
            for item in sorted(
                competencia.conocimientos,
                key=lambda item: (item.orden is None, item.orden or 0),
            )
        ],
        "criterios": [
            {
                "id": str(item.id),
                "competencia_id": str(item.competencia_id),
                "resultado_id": str(item.resultado_id) if item.resultado_id else None,
                "descripcion": item.descripcion,
                "orden": item.orden,
                "estado": item.estado.value,
                "fecha_creacion": item.fecha_creacion.isoformat(),
                "fecha_actualizacion": item.fecha_actualizacion.isoformat(),
            }
            for item in sorted(
                competencia.criterios,
                key=lambda item: (item.orden is None, item.orden or 0),
            )
        ],
    }


def _required_string(
    row: dict[str, object],
    sheet_name: str,
    field: str,
    errores: list[ExcelValidationIssueDTO],
) -> str | None:
    value = _optional_string(row.get(field))
    if value is None:
        errores.append(
            ExcelValidationIssueDTO(
                hoja=sheet_name,
                fila=_row_index(row),
                campo=field,
                mensaje=f"{field} es obligatorio",
            ),
        )
        return None
    return value


def _optional_int(
    row: dict[str, object],
    sheet_name: str,
    field: str,
    errores: list[ExcelValidationIssueDTO],
) -> int | None:
    value = row.get(field)
    if value is None or value == "":
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and re.fullmatch(r"\d+", value.strip()):
        return int(value.strip())
    errores.append(
        ExcelValidationIssueDTO(
            hoja=sheet_name,
            fila=_row_index(row),
            campo=field,
            mensaje=f"{field} debe ser numerico entero cuando se informa",
        ),
    )
    return None


def _parse_tipo_conocimiento(
    value: str | None,
    row: dict[str, object],
    errores: list[ExcelValidationIssueDTO],
) -> TipoConocimiento | None:
    if value is None:
        return None
    normalized = value.strip().upper()
    try:
        return TipoConocimiento(normalized)
    except ValueError:
        errores.append(
            ExcelValidationIssueDTO(
                hoja="Conocimientos",
                fila=_row_index(row),
                campo="tipo_conocimiento",
                mensaje="tipo_conocimiento solo puede ser SABER o PROCESO",
            ),
        )
        return None


def _optional_string(value: object) -> str | None:
    if value is None:
        return None
    clean = str(value).strip()
    return clean or None


def _cell_to_string(value: object) -> str:
    return "" if value is None else str(value).strip()


def _public_record(row: dict[str, object]) -> dict[str, object]:
    public = {key: value for key, value in row.items() if not key.startswith("_")}
    public["_row_index"] = row.get("_row_index")
    return public


def _row_index(row: dict[str, object]) -> int | None:
    value = row.get("_row_index")
    return value if isinstance(value, int) else None


def _norm(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().lower()


def _as_record(value: object) -> dict[str, object] | None:
    return value if isinstance(value, dict) else None


def _read_string(record: dict[str, object], key: str) -> str:
    value = record.get(key)
    return str(value) if value is not None else ""


def _append_touched_step(value: object, step: str) -> list[str]:
    if isinstance(value, list):
        steps = [item for item in value if isinstance(item, str)]
        if step not in steps:
            steps.append(step)
        return steps
    return [step]


def _build_excel_storage_key(referencia_id: uuid.UUID, filename: str) -> str:
    safe_filename = re.sub(r"[^a-zA-Z0-9._-]+", "-", filename).strip("-")
    if not safe_filename:
        safe_filename = "programa.xlsx"
    object_id = uuid.uuid4()
    return f"programas/{referencia_id}/documentos/{object_id}-{safe_filename}"
