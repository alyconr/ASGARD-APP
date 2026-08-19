"""Generate the official GPFI-F-134 V05 planning workbook."""

from __future__ import annotations

import copy
import hashlib
import io
import math
import warnings
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import cast
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

EXCEL_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
OFFICIAL_FORMAT_CODE = "GPFI-F-134"
OFFICIAL_FORMAT_VERSION = "05"
OFFICIAL_FILE_NAME = "GPFI-F-134V05-planeacion-pedagogica.xlsx"
TEMPLATE_PATH = (
    Path(__file__).resolve().parents[2]
    / "infrastructure"
    / "templates"
    / "planeacion"
    / "GPFI-F-134V05.xlsx"
)

_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_DOC_REL_NS = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
)
_PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_XDR_NS = (
    "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
)
_DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


class PlaneacionFormatoValidationError(ValueError):
    """Raised when official workbook generation has domain gaps."""

    def __init__(self, messages: list[str]) -> None:
        self.messages = messages
        super().__init__("; ".join(messages))


class PlaneacionTemplateMissingError(FileNotFoundError):
    """Raised when the canonical institutional workbook is unavailable."""


@dataclass(frozen=True)
class FormatoPlaneacionMetadata:
    """General header values shared by all generated rows."""

    fecha_elaboracion: date
    nombre_programa: str
    modalidad_formacion: str
    codigo_programa: str
    version_programa: str | None
    nombre_proyecto: str
    codigo_proyecto: str
    equipo_gestion_curricular: tuple[str, ...]
    regional: str
    centro_formacion: str
    clasificacion_informacion: str


@dataclass(frozen=True)
class FormatoPlaneacionRow:
    """One official FASE row.

    Hours belong to the integrated learning activity, not to each RAP.
    When a planning block expands into several rows, only the first row
    carries the hour values; the rest keep ``None`` so the official sheet
    does not multiply the activity duration.
    """

    fase: str
    actividad_proyecto: str
    competencia: str
    resultado: str
    saberes: tuple[str, ...]
    procesos: tuple[str, ...]
    criterios: tuple[str, ...]
    actividades_aprendizaje: str
    horas_trabajo_directo: float | None
    horas_trabajo_independiente: float | None
    descripcion_evidencia: str
    estrategias_didacticas: str
    ambiente: tuple[str, ...]
    materiales_formacion: str
    instructores: str
    observaciones: str


@dataclass(frozen=True)
class FormatoExcelResultado:
    """Generated workbook bytes and integrity metadata."""

    content: bytes
    checksum_sha256: str
    filas_generadas: int


class PlaneacionFormatoExcelService:
    """Fill a fresh copy of the official institutional workbook."""

    def __init__(self, template_path: Path = TEMPLATE_PATH) -> None:
        self._template_path = template_path

    def generar(
        self,
        *,
        metadata: FormatoPlaneacionMetadata,
        rows: list[FormatoPlaneacionRow],
    ) -> FormatoExcelResultado:
        """Return a valid official workbook without mutating the template."""
        self._validate(metadata=metadata, rows=rows)
        if not self._template_path.is_file():
            raise PlaneacionTemplateMissingError(
                f"No existe la plantilla oficial {self._template_path}"
            )

        template_bytes = self._template_path.read_bytes()
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            workbook = load_workbook(io.BytesIO(template_bytes))

        if workbook.sheetnames != ["Instrucciones", "FASE"]:
            raise PlaneacionTemplateMissingError(
                "La plantilla oficial debe contener Instrucciones y FASE en ese orden"
            )

        sheet = workbook["FASE"]
        data_start = self._find_data_start(sheet)
        self._fill_header(sheet, metadata)
        self._fill_rows(sheet, data_start=data_start, rows=rows)
        last_row = data_start + len(rows) - 1
        sheet.print_area = f"A1:P{last_row}"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_LEGAL
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        if sheet.sheet_properties.pageSetUpPr is None:
            sheet.sheet_properties.pageSetUpPr = PageSetupProperties()
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

        generated = io.BytesIO()
        workbook.save(generated)
        content = self._merge_generated_sheet_into_template(
            template_bytes=template_bytes,
            generated_bytes=generated.getvalue(),
            print_last_row=last_row,
            classification=metadata.clasificacion_informacion,
        )
        return FormatoExcelResultado(
            content=content,
            checksum_sha256=hashlib.sha256(content).hexdigest(),
            filas_generadas=len(rows),
        )

    @staticmethod
    def _validate(
        *,
        metadata: FormatoPlaneacionMetadata,
        rows: list[FormatoPlaneacionRow],
    ) -> None:
        missing: list[str] = []
        required = {
            "fecha de elaboracion": metadata.fecha_elaboracion,
            "programa de formacion": metadata.nombre_programa.strip(),
            "modalidad de formacion": metadata.modalidad_formacion.strip(),
            "codigo del programa": metadata.codigo_programa.strip(),
            "proyecto formativo": metadata.nombre_proyecto.strip(),
            "codigo del proyecto": metadata.codigo_proyecto.strip(),
            "equipo de gestion curricular": metadata.equipo_gestion_curricular,
            "regional": metadata.regional.strip(),
            "centro de formacion": metadata.centro_formacion.strip(),
            "clasificacion de la informacion": (
                metadata.clasificacion_informacion.strip()
            ),
        }
        missing.extend(name for name, value in required.items() if not value)
        if metadata.clasificacion_informacion not in {
            "PUBLICA",
            "PUBLICA_CLASIFICADA",
            "PUBLICA_RESERVADA",
        }:
            missing.append("clasificacion de la informacion invalida")
        if not rows:
            missing.append("al menos una fila exportable")
        for index, row in enumerate(rows, start=1):
            horas = (row.horas_trabajo_directo, row.horas_trabajo_independiente)
            if any(value is not None and value < 0 for value in horas):
                missing.append(f"fila {index}: las horas no pueden ser negativas")
        if missing:
            raise PlaneacionFormatoValidationError(missing)

    @staticmethod
    def _find_data_start(sheet: Worksheet) -> int:
        header_row = next(
            (
                row
                for row in range(1, sheet.max_row + 1)
                if str(sheet.cell(row, 1).value or "").strip().upper().startswith(
                    "FASE DE PROYECTO FORMATIVO"
                )
            ),
            None,
        )
        if header_row is None:
            raise PlaneacionTemplateMissingError(
                "No se encontro el encabezado de la tabla oficial en FASE"
            )
        header_end = max(
            (
                merged.max_row
                for merged in sheet.merged_cells.ranges
                if merged.min_row <= header_row <= merged.max_row
            ),
            default=header_row,
        )
        return header_end + 1

    @staticmethod
    def _fill_header(
        sheet: Worksheet,
        metadata: FormatoPlaneacionMetadata,
    ) -> None:
        values: dict[str, date | str] = {
            "E9": metadata.fecha_elaboracion,
            "E10": metadata.nombre_programa,
            "E11": metadata.modalidad_formacion,
            "E12": " / ".join(
                value
                for value in (
                    metadata.codigo_programa,
                    metadata.version_programa,
                )
                if value
            ),
            "E13": metadata.nombre_proyecto,
            "E14": metadata.codigo_proyecto,
            "E15": "\n".join(metadata.equipo_gestion_curricular),
            "K15": f"{metadata.regional}\n{metadata.centro_formacion}",
        }
        for coordinate, value in values.items():
            cell = cast(Cell, sheet[coordinate])
            cell.value = value
            alignment = copy.copy(cell.alignment)
            cell.alignment = Alignment(
                horizontal=alignment.horizontal,
                vertical="center",
                text_rotation=alignment.text_rotation,
                wrap_text=True,
                shrink_to_fit=alignment.shrink_to_fit,
                indent=alignment.indent,
            )
        sheet["E9"].number_format = "dd/mm/yyyy"

    def _fill_rows(
        self,
        sheet: Worksheet,
        *,
        data_start: int,
        rows: list[FormatoPlaneacionRow],
    ) -> None:
        pattern_row = data_start
        required_last_row = data_start + len(rows) - 1
        if required_last_row > sheet.max_row:
            sheet.insert_rows(
                sheet.max_row + 1,
                amount=required_last_row - sheet.max_row,
            )

        for row_index, row in enumerate(rows, start=data_start):
            self._copy_row_format(sheet, pattern_row, row_index)
            values: tuple[str | float | None, ...] = (
                row.fase,
                row.actividad_proyecto,
                row.competencia,
                row.resultado,
                self._as_lines(row.saberes),
                self._as_lines(row.procesos),
                self._as_lines(row.criterios),
                row.actividades_aprendizaje,
                row.horas_trabajo_directo,
                row.horas_trabajo_independiente,
                row.descripcion_evidencia,
                row.estrategias_didacticas,
                self._as_lines(row.ambiente),
                row.materiales_formacion,
                row.instructores,
                row.observaciones,
            )
            for column, value in enumerate(values, start=1):
                sheet.cell(row_index, column, value)
            sheet.row_dimensions[row_index].height = self._row_height(sheet, row_index)

    @staticmethod
    def _copy_row_format(
        sheet: Worksheet,
        source_row: int,
        target_row: int,
    ) -> None:
        source_height = sheet.row_dimensions[source_row].height
        if source_height is not None:
            sheet.row_dimensions[target_row].height = source_height
        for column in range(1, 17):
            source = cast(Cell, sheet.cell(source_row, column))
            target = cast(Cell, sheet.cell(target_row, column))
            setattr(target, "_style", copy.copy(getattr(source, "_style")))
            setattr(target, "font", copy.copy(source.font))
            setattr(target, "fill", copy.copy(source.fill))
            setattr(target, "border", copy.copy(source.border))
            setattr(target, "alignment", copy.copy(source.alignment))
            target.number_format = source.number_format
            setattr(target, "protection", copy.copy(source.protection))

    @staticmethod
    def _as_lines(values: tuple[str, ...]) -> str:
        return "\n".join(f"- {value}" for value in values)

    @staticmethod
    def _row_height(sheet: Worksheet, row: int) -> float:
        max_lines = 1
        for column in range(1, 17):
            value = str(sheet.cell(row, column).value or "")
            width = sheet.column_dimensions[get_column_letter(column)].width
            chars_per_line = max(8, int((width or 10) * 1.35))
            line_count = sum(
                max(1, math.ceil(len(part) / chars_per_line))
                for part in value.splitlines() or [""]
            )
            max_lines = max(max_lines, line_count)
        return min(180.0, max(30.0, 13.5 * max_lines + 8))

    def _merge_generated_sheet_into_template(
        self,
        *,
        template_bytes: bytes,
        generated_bytes: bytes,
        print_last_row: int,
        classification: str,
    ) -> bytes:
        with ZipFile(io.BytesIO(template_bytes)) as template_zip:
            with ZipFile(io.BytesIO(generated_bytes)) as generated_zip:
                sheet_part, sheet_index = self._worksheet_part(
                    template_zip,
                    "FASE",
                )
                generated_sheet_part, _ = self._worksheet_part(
                    generated_zip,
                    "FASE",
                )
                sheet_xml = self._restore_sheet_relationships(
                    generated=generated_zip.read(generated_sheet_part),
                    original=template_zip.read(sheet_part),
                )
                workbook_xml = self._set_print_area(
                    template_zip.read("xl/workbook.xml"),
                    sheet_index=sheet_index,
                    last_row=print_last_row,
                )
                drawing_part = self._drawing_part_for_sheet(
                    template_zip,
                    sheet_part,
                )
                drawing_xml = self._mark_classification(
                    template_zip.read(drawing_part),
                    classification,
                )

                output = io.BytesIO()
                with ZipFile(output, "w", ZIP_DEFLATED) as result_zip:
                    for info in template_zip.infolist():
                        content = template_zip.read(info.filename)
                        if info.filename == sheet_part:
                            content = sheet_xml
                        elif info.filename == "xl/workbook.xml":
                            content = workbook_xml
                        elif info.filename == "xl/styles.xml":
                            content = generated_zip.read("xl/styles.xml")
                        elif info.filename == drawing_part:
                            content = drawing_xml
                        result_zip.writestr(info, content)
                return output.getvalue()

    @staticmethod
    def _worksheet_part(workbook_zip: ZipFile, sheet_name: str) -> tuple[str, int]:
        workbook = ET.fromstring(workbook_zip.read("xl/workbook.xml"))
        relations = ET.fromstring(
            workbook_zip.read("xl/_rels/workbook.xml.rels")
        )
        targets = {
            relation.attrib["Id"]: relation.attrib["Target"].lstrip("/")
            for relation in relations.findall(f"{{{_PKG_REL_NS}}}Relationship")
        }
        sheets = workbook.find(f"{{{_MAIN_NS}}}sheets")
        if sheets is None:
            raise PlaneacionTemplateMissingError("La plantilla no contiene hojas")
        for index, sheet in enumerate(sheets):
            if sheet.attrib.get("name") != sheet_name:
                continue
            relation_id = sheet.attrib.get(f"{{{_DOC_REL_NS}}}id")
            if not relation_id or relation_id not in targets:
                break
            target = targets[relation_id]
            return (
                target if target.startswith("xl/") else f"xl/{target}",
                index,
            )
        raise PlaneacionTemplateMissingError(
            f"No se encontro la hoja {sheet_name} en la plantilla"
        )

    @staticmethod
    def _drawing_part_for_sheet(workbook_zip: ZipFile, sheet_part: str) -> str:
        relation_part = (
            f"{sheet_part.rsplit('/', 1)[0]}/_rels/"
            f"{sheet_part.rsplit('/', 1)[1]}.rels"
        )
        relations = ET.fromstring(workbook_zip.read(relation_part))
        for relation in relations.findall(f"{{{_PKG_REL_NS}}}Relationship"):
            if relation.attrib.get("Type", "").endswith("/drawing"):
                target = relation.attrib["Target"]
                return f"xl/{target.removeprefix('../')}"
        raise PlaneacionTemplateMissingError(
            "La hoja FASE no conserva su DrawingML institucional"
        )

    @staticmethod
    def _register_namespaces(xml: bytes) -> None:
        for _, namespace in ET.iterparse(io.BytesIO(xml), events=("start-ns",)):
            prefix, uri = cast(tuple[str, str], namespace)
            if prefix not in {"xml", "xmlns"}:
                ET.register_namespace(prefix or "", uri)

    @staticmethod
    def _restore_sheet_relationships(*, generated: bytes, original: bytes) -> bytes:
        PlaneacionFormatoExcelService._register_namespaces(original)
        PlaneacionFormatoExcelService._register_namespaces(generated)
        generated_root = ET.fromstring(generated)
        original_root = ET.fromstring(original)
        relationship_elements = {"drawing", "legacyDrawing", "legacyDrawingHF"}
        for element in list(generated_root):
            if element.tag.split("}")[-1] in relationship_elements:
                generated_root.remove(element)
        for element in original_root:
            if element.tag.split("}")[-1] in relationship_elements:
                generated_root.append(copy.deepcopy(element))

        original_page_setup = original_root.find(f"{{{_MAIN_NS}}}pageSetup")
        generated_page_setup = generated_root.find(f"{{{_MAIN_NS}}}pageSetup")
        if original_page_setup is not None and generated_page_setup is not None:
            relation_id = original_page_setup.attrib.get(f"{{{_DOC_REL_NS}}}id")
            if relation_id:
                generated_page_setup.set(f"{{{_DOC_REL_NS}}}id", relation_id)
        return cast(
            bytes,
            ET.tostring(
                generated_root,
                encoding="utf-8",
                xml_declaration=True,
            ),
        )

    @staticmethod
    def _set_print_area(
        workbook_xml: bytes,
        *,
        sheet_index: int,
        last_row: int,
    ) -> bytes:
        PlaneacionFormatoExcelService._register_namespaces(workbook_xml)
        root = ET.fromstring(workbook_xml)
        defined_names = root.find(f"{{{_MAIN_NS}}}definedNames")
        if defined_names is None:
            defined_names = ET.Element(f"{{{_MAIN_NS}}}definedNames")
            sheets = root.find(f"{{{_MAIN_NS}}}sheets")
            insert_at = (
                list(root).index(sheets) + 1
                if sheets is not None
                else len(root)
            )
            root.insert(insert_at, defined_names)
        for item in list(defined_names):
            if (
                item.attrib.get("name") == "_xlnm.Print_Area"
                and item.attrib.get("localSheetId") == str(sheet_index)
            ):
                defined_names.remove(item)
        print_area = ET.SubElement(
            defined_names,
            f"{{{_MAIN_NS}}}definedName",
            {"name": "_xlnm.Print_Area", "localSheetId": str(sheet_index)},
        )
        print_area.text = f"'FASE'!$A$1:$P${last_row}"
        return cast(
            bytes,
            ET.tostring(root, encoding="utf-8", xml_declaration=True),
        )

    @staticmethod
    def _mark_classification(drawing_xml: bytes, classification: str) -> bytes:
        PlaneacionFormatoExcelService._register_namespaces(drawing_xml)
        ET.register_namespace("xdr", _XDR_NS)
        ET.register_namespace("a", _DRAWING_NS)
        root = ET.fromstring(drawing_xml)
        anchors = list(root)
        public_anchor = next(
            (
                anchor
                for anchor in anchors
                if anchor.find(f"{{{_XDR_NS}}}sp") is not None
                and anchor.findtext(
                    f"{{{_XDR_NS}}}from/{{{_XDR_NS}}}row"
                )
                == "7"
            ),
            None,
        )
        if public_anchor is None:
            raise PlaneacionTemplateMissingError(
                "No se encontro el marcador de clasificacion institucional"
            )

        target_column = {
            "PUBLICA": 3,
            "PUBLICA_CLASIFICADA": 7,
            "PUBLICA_RESERVADA": 13,
        }[classification]
        target_anchor = public_anchor
        if target_column != 3:
            target_anchor = copy.deepcopy(public_anchor)
            delta = target_column - 3
            for marker in ("from", "to"):
                column = target_anchor.find(
                    f"{{{_XDR_NS}}}{marker}/{{{_XDR_NS}}}col"
                )
                if column is not None and column.text is not None:
                    column.text = str(int(column.text) + delta)
            root.append(target_anchor)

        shape = target_anchor.find(f"{{{_XDR_NS}}}sp")
        if shape is None:
            raise PlaneacionTemplateMissingError(
                "El marcador de clasificacion no es editable"
            )
        paragraph = shape.find(
            f"{{{_XDR_NS}}}txBody/{{{_DRAWING_NS}}}p"
        )
        if paragraph is None:
            raise PlaneacionTemplateMissingError(
                "El marcador de clasificacion no admite texto"
            )
        run = ET.Element(f"{{{_DRAWING_NS}}}r")
        run_properties = ET.SubElement(
            run,
            f"{{{_DRAWING_NS}}}rPr",
            {"lang": "es-CO", "sz": "900", "b": "1"},
        )
        ET.SubElement(
            run_properties,
            f"{{{_DRAWING_NS}}}solidFill",
        ).append(ET.Element(f"{{{_DRAWING_NS}}}srgbClr", {"val": "000000"}))
        ET.SubElement(run, f"{{{_DRAWING_NS}}}t").text = "X"
        paragraph.insert(max(0, len(paragraph) - 1), run)
        return cast(
            bytes,
            ET.tostring(root, encoding="utf-8", xml_declaration=True),
        )
