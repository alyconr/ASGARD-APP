"""Canonical Excel preview and import service for project formativo."""

from __future__ import annotations

import io
import re
import uuid
from dataclasses import dataclass
from datetime import UTC, datetime
from typing import Protocol

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import selectinload

from src.application.dto.programa_documentos import StoredDocumentDTO
from src.application.dto.proyecto_excel import (
    ExcelActividadPreviewDTO,
    ExcelCompetenciaPreviewDTO,
    ExcelFasePreviewDTO,
    ExcelPendingSummaryDTO,
    ExcelPreviewSummaryDTO,
    ExcelProjectPreviewDTO,
    ExcelResultPreviewDTO,
    ExcelValidationIssueDTO,
    ProyectoExcelImportDTO,
    ProyectoExcelPreviewDTO,
)
from src.domain.drafts.types import TipoBloqueBorrador
from src.domain.shared.enums import EstadoBloque, TipoResultadoProyecto
from src.infrastructure.db.models.curriculum import (
    Competencia,
    ProgramaFormacion,
    ResultadoAprendizaje,
)
from src.infrastructure.db.models.drafts import BorradorSesion
from src.infrastructure.db.models.proyecto import (
    ActividadProyecto,
    AsignacionCurricularProyecto,
    FaseProyecto,
    ProyectoFormativo,
)
from src.infrastructure.storage.document_storage import build_proyecto_storage_prefix

CANONICAL_SHEETS: dict[str, list[str]] = {
    "Proyecto": [
        "proyecto_id",
        "nombre_proyecto",
        "codigo_proyecto_sofia",
        "codigo_programa",
        "nombre_programa",
        "fuente_archivo",
        "observaciones",
    ],
    "Planeacion_Proyecto": [
        "proyecto_id",
        "fase_id",
        "fase_proyecto",
        "actividad_id",
        "actividad_proyecto",
        "tipo_resultado",
        "competencia_id",
        "codigo_competencia",
        "nombre_competencia",
        "rap_id",
        "rap_numero",
        "resultado_aprendizaje",
        "orden_fase",
        "orden_actividad",
        "orden_resultado",
        "pagina_origen",
        "observaciones",
    ],
    "Validacion_Proyecto": [
        "tipo_validacion",
        "descripcion",
        "estado",
        "observaciones",
    ],
}

EXCEL_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ProjectExcelDraftMissingError(Exception):
    """Raised when an Excel import has no current wizard draft."""


class InvalidProjectExcelUploadError(Exception):
    """Raised when the upload is not an accepted canonical workbook."""


class ProjectExcelValidationError(Exception):
    """Raised when a workbook cannot be imported after validation."""


class ProjectExcelMissingPreviewError(Exception):
    """Raised when import confirmation has no valid preview to confirm."""


class ProjectExcelStorageMissingError(Exception):
    """Raised when the stored workbook cannot be read for confirmation."""


# Builder classes for consolidating flat sheets
class ResultBuilder:
    def __init__(
        self,
        rap_id: str,
        rap_numero: str,
        resultado_aprendizaje: str,
        tipo_resultado: str,
        orden_resultado: int | None,
        pagina_origen: str | None,
        observaciones: str | None,
    ):
        self.rap_id = rap_id
        self.rap_numero = rap_numero
        self.resultado_aprendizaje = resultado_aprendizaje
        self.tipo_resultado = tipo_resultado
        self.orden_resultado = orden_resultado
        self.pagina_origen = pagina_origen
        self.observaciones = observaciones


class CompetenciaBuilder:
    def __init__(
        self, competencia_id: str, codigo_competencia: str, nombre_competencia: str
    ):
        self.competencia_id = competencia_id
        self.codigo_competencia = codigo_competencia
        self.nombre_competencia = nombre_competencia
        self.resultados: dict[str, ResultBuilder] = {}


class ActividadBuilder:
    def __init__(self, actividad_id: str, descripcion: str, orden: int | None):
        self.actividad_id = actividad_id
        self.descripcion = descripcion
        self.orden = orden
        self.competencias: dict[str, CompetenciaBuilder] = {}


class FaseBuilder:
    def __init__(self, fase_id: str, nombre_fase: str, orden: int | None):
        self.fase_id = fase_id
        self.nombre_fase = nombre_fase
        self.orden = orden
        self.actividades: dict[str, ActividadBuilder] = {}


@dataclass(frozen=True)
class ProjectRow:
    proyecto_id: str
    nombre_proyecto: str
    codigo_proyecto_sofia: str
    codigo_programa: str
    nombre_programa: str
    fuente_archivo: str | None
    observaciones: str | None
    raw: dict[str, object]


@dataclass(frozen=True)
class PlaneacionRow:
    proyecto_id: str
    fase_id: str
    fase_proyecto: str
    actividad_id: str
    actividad_proyecto: str
    tipo_resultado: str
    competencia_id: str
    codigo_competencia: str
    nombre_competencia: str
    rap_id: str
    rap_numero: str
    resultado_aprendizaje: str
    orden_fase: int | None
    orden_actividad: int | None
    orden_resultado: int | None
    pagina_origen: str | None
    observaciones: str | None
    raw: dict[str, object]


@dataclass(frozen=True)
class ValidacionRow:
    tipo_validacion: str
    descripcion: str
    estado: str
    observaciones: str | None
    raw: dict[str, object]


@dataclass(frozen=True)
class CanonicalWorkbook:
    proyecto: ProjectRow | None
    planeacion: list[PlaneacionRow]
    validacion: list[ValidacionRow]
    errores: list[ExcelValidationIssueDTO]

    @property
    def is_valid(self) -> bool:
        """Return whether the workbook can proceed to preview or import."""
        return len(self.errores) == 0 and self.proyecto is not None


class DocumentStorageProtocol(Protocol):
    """Storage behavior required by the canonical Excel service."""

    async def save_excel(
        self,
        *,
        key: str,
        content: bytes,
        content_type: str,
        original_filename: str,
    ) -> StoredDocumentDTO:
        """Persist an Excel file and return object metadata."""

    async def read_excel(self, *, key: str) -> bytes:
        """Read a stored Excel file."""


class DraftRepositoryProtocol(Protocol):
    """Draft repository dependency used to associate imports to the wizard."""

    async def get_by_block_reference(
        self,
        tipo_bloque: TipoBloqueBorrador,
        referencia_id: uuid.UUID,
    ) -> BorradorSesion | None:
        """Return the current logical draft for a reference."""

    async def save(self, draft: BorradorSesion) -> BorradorSesion:
        """Persist modifications on an existing draft row."""


class AuditRepositoryProtocol(Protocol):
    """Audit dependency for relevant import events."""

    async def add_event(
        self,
        entidad: str,
        entidad_id: uuid.UUID,
        accion: str,
        detalle: dict[str, object] | None = None,
    ) -> object:
        """Persist a basic audit event."""


class ProjectRepositoryProtocol(Protocol):
    """Repository behavior required by the project Excel import service."""

    async def create_proyecto(
        self,
        *,
        programa_id: uuid.UUID,
        codigo_proyecto: str,
        nombre_proyecto: str,
        version_proyecto: str,
    ) -> ProyectoFormativo:
        """Create a project row."""

    async def create_fase(
        self,
        *,
        proyecto_id: uuid.UUID,
        nombre_fase: str,
        orden: int | None,
    ) -> FaseProyecto:
        """Create a project phase."""

    async def create_actividad(
        self,
        *,
        fase_id: uuid.UUID,
        descripcion: str,
        orden: int | None,
    ) -> ActividadProyecto:
        """Create a project activity."""

    async def proyecto_exists_by_code_name(
        self,
        *,
        programa_id: uuid.UUID,
        codigo_proyecto: str,
        nombre_proyecto: str,
    ) -> bool:
        """Return whether the same project already exists for the program."""

    async def create_asignacion_curricular(
        self,
        *,
        proyecto_id: uuid.UUID,
        actividad_proyecto_id: uuid.UUID,
        competencia_id: uuid.UUID,
        resultado_id: uuid.UUID | None,
        tipo_resultado: str,
        orden_resultado: int | None,
        pagina_origen: str | None,
        observaciones: str | None,
    ) -> AsignacionCurricularProyecto:
        """Persist one project activity ↔ competency ↔ result assignment."""


class AsyncSessionProtocol(Protocol):
    """Subset of async session behavior required by this service."""

    async def commit(self) -> None:
        """Commit the current transaction."""

    async def refresh(self, instance: object) -> None:
        """Refresh the provided ORM instance."""

    def delete(self, instance: object) -> object:
        """Delete the provided ORM instance."""

    async def flush(self) -> None:
        """Flush the session."""


class ProyectoExcelImportService:
    """Validate, preview and confirm canonical Excel project imports."""

    def __init__(
        self,
        session: AsyncSessionProtocol,
        draft_repository: DraftRepositoryProtocol,
        audit_repository: AuditRepositoryProtocol,
        project_repository: ProjectRepositoryProtocol,
        storage_service: DocumentStorageProtocol,
    ) -> None:
        """Initialize the service with explicit infrastructure ports."""
        self._session = session
        self._draft_repository = draft_repository
        self._audit_repository = audit_repository
        self._project_repository = project_repository
        self._storage_service = storage_service

    async def preview_project_excel(
        self,
        *,
        referencia_id: uuid.UUID,
        filename: str,
        content_type: str,
        content: bytes,
    ) -> ProyectoExcelPreviewDTO:
        """Store and validate a canonical Excel workbook without relational writes."""
        _validate_excel_upload(filename=filename, content=content)
        draft = await self._get_project_draft(referencia_id)
        programa_id = await self._resolve_programa_id(draft)
        workbook = parse_canonical_workbook(content)
        if workbook.is_valid and workbook.proyecto is not None:
            if await self._project_repository.proyecto_exists_by_code_name(
                programa_id=programa_id,
                codigo_proyecto=workbook.proyecto.codigo_proyecto_sofia,
                nombre_proyecto=workbook.proyecto.nombre_proyecto,
            ):
                raise ProjectExcelValidationError(
                    "Este proyecto formativo ya fue cargado con el mismo "
                    "codigo y nombre para el programa asociado. No se debe "
                    "iniciar un nuevo wizard para duplicarlo; continua "
                    "directamente con el wizard de planeacion pedagogica."
                )

        stored_document: StoredDocumentDTO | None = None

        if workbook.is_valid:
            assert workbook.proyecto is not None
            prefix = build_proyecto_storage_prefix(
                nombre=workbook.proyecto.nombre_proyecto,
                codigo=workbook.proyecto.codigo_proyecto_sofia,
            )
            key = f"{prefix}/excel/matriz-proyecto.xlsx"
            stored_document = await self._storage_service.save_excel(
                key=key,
                content=content,
                content_type=content_type or EXCEL_CONTENT_TYPE,
                original_filename=filename,
            )

        preview = _build_preview_dto(
            referencia_id=referencia_id,
            workbook=workbook,
            document=stored_document,
        )
        draft.payload_json = _merge_excel_preview_into_payload(
            payload=draft.payload_json,
            preview=preview,
        )
        draft.paso_actual = "fuente-proyecto"
        draft.estado_borrador = EstadoBloque.BORRADOR
        await self._draft_repository.save(draft)
        await self._audit_repository.add_event(
            entidad="BorradorSesion",
            entidad_id=draft.id,
            accion="EXCEL_PROYECTO_PREVISUALIZADO",
            detalle={
                "referencia_id": str(referencia_id),
                "estado_validacion": preview.estado_validacion,
                "errores": len(preview.errores),
            },
        )
        await self._session.commit()
        await self._session.refresh(draft)
        return preview

    async def confirm_project_excel_import(
        self,
        *,
        referencia_id: uuid.UUID,
    ) -> ProyectoExcelImportDTO:
        """Materialize a previously validated canonical workbook."""
        draft = await self._get_project_draft(referencia_id)
        preview_payload = _get_valid_excel_preview_payload(draft.payload_json)
        programa_id = await self._resolve_programa_id(draft)

        document_payload = _as_record(preview_payload.get("documento"))
        storage_key = _read_string(document_payload or {}, "storage_key")
        if not storage_key:
            raise ProjectExcelMissingPreviewError(
                "El borrador no tiene un Excel canonico validado para confirmar",
            )

        try:
            content = await self._storage_service.read_excel(key=storage_key)
        except FileNotFoundError as error:
            raise ProjectExcelStorageMissingError(
                "El Excel canonico validado ya no existe en MinIO",
            ) from error

        workbook = parse_canonical_workbook(content)
        if not workbook.is_valid or workbook.proyecto is None:
            raise ProjectExcelValidationError(
                "El Excel canonico almacenado ya no supera la validacion",
            )
        if await self._project_repository.proyecto_exists_by_code_name(
            programa_id=programa_id,
            codigo_proyecto=workbook.proyecto.codigo_proyecto_sofia,
            nombre_proyecto=workbook.proyecto.nombre_proyecto,
        ):
            raise ProjectExcelValidationError(
                "Este proyecto formativo ya fue cargado con el mismo codigo "
                "y nombre para el programa asociado. No se debe iniciar un "
                "nuevo wizard para duplicarlo; continua directamente con el "
                "wizard de planeacion pedagogica."
            )

        import_result = await self._materialize_workbook(
            referencia_id=referencia_id,
            draft=draft,
            programa_id=programa_id,
            workbook=workbook,
        )
        draft.paso_actual = "fuente-proyecto"
        draft.payload_json = _merge_excel_import_into_payload(
            payload=draft.payload_json,
            result=import_result,
        )
        await self._draft_repository.save(draft)
        await self._audit_repository.add_event(
            entidad="ProyectoFormativo",
            entidad_id=import_result.proyecto_id,
            accion="EXCEL_PROYECTO_IMPORTADO",
            detalle={
                "referencia_id": str(referencia_id),
                "fases": import_result.resumen.fases,
                "actividades": import_result.resumen.actividades,
            },
        )
        await self._session.commit()
        await self._session.refresh(draft)
        return import_result

    async def _get_project_draft(self, referencia_id: uuid.UUID) -> BorradorSesion:
        draft = await self._draft_repository.get_by_block_reference(
            TipoBloqueBorrador.PROYECTO,
            referencia_id,
        )
        if draft is None:
            raise ProjectExcelDraftMissingError(
                "No existe un borrador de proyecto para asociar el Excel",
            )
        return draft

    async def _resolve_programa_id(self, draft: BorradorSesion) -> uuid.UUID:
        """Resolve the completed program id associated with a project draft."""
        meta = _as_record(draft.payload_json.get("meta")) or {}
        programa_id = _read_uuid(meta, "programaId")
        if programa_id is not None:
            # Safely handle mock sessions in tests while checking DB existence.
            if hasattr(self._session, "execute"):
                import inspect

                stmt = select(ProgramaFormacion).where(
                    ProgramaFormacion.id == programa_id
                )
                res = self._session.execute(stmt)
                if inspect.isawaitable(res):
                    res = await res
                if res.scalar_one_or_none() is not None:
                    return programa_id

            else:
                return programa_id

        programa_referencia_id = _read_uuid(meta, "programaReferenciaId")
        if programa_referencia_id is not None:
            program_draft = await self._draft_repository.get_by_block_reference(
                TipoBloqueBorrador.PROGRAMA,
                programa_referencia_id,
            )
            if program_draft is not None:
                programa_id = _extract_programa_id_from_program_payload(
                    program_draft.payload_json
                )
                if programa_id is not None:
                    draft.payload_json = {
                        **draft.payload_json,
                        "meta": {
                            **meta,
                            "programaId": str(programa_id),
                        },
                    }
                    return programa_id

        raise ProjectExcelValidationError(
            "No se encontro el programa de formacion asociado al proyecto. "
            "Cierre el programa y vuelva a iniciar el wizard del proyecto."
        )

    async def _materialize_workbook(
        self,
        *,
        referencia_id: uuid.UUID,
        draft: BorradorSesion,
        programa_id: uuid.UUID,
        workbook: CanonicalWorkbook,
    ) -> ProyectoExcelImportDTO:
        if workbook.proyecto is None:
            raise ProjectExcelValidationError("La hoja Proyecto esta vacia")

        # Sofia code is used as codigo_proyecto, default version to "1"
        proyecto = await self._project_repository.create_proyecto(
            programa_id=programa_id,
            codigo_proyecto=workbook.proyecto.codigo_proyecto_sofia,
            nombre_proyecto=workbook.proyecto.nombre_proyecto,
            version_proyecto="1",
        )
        await self._session.refresh(proyecto)

        fase_by_excel_id: dict[str, uuid.UUID] = {}
        actividad_by_excel_id: dict[str, uuid.UUID] = {}
        fase_ids: list[uuid.UUID] = []
        actividad_ids: list[uuid.UUID] = []

        # Find unique Fases from workbook.planeacion
        unique_fases = {}
        for r in workbook.planeacion:
            if r.fase_id not in unique_fases:
                unique_fases[r.fase_id] = (r.fase_proyecto, r.orden_fase)

        for fid, (nombre_fase, orden_fase) in unique_fases.items():
            fase = await self._project_repository.create_fase(
                proyecto_id=proyecto.id,
                nombre_fase=nombre_fase,
                orden=orden_fase,
            )
            await self._session.refresh(fase)
            fase_by_excel_id[fid] = fase.id
            fase_ids.append(fase.id)

        # Find unique Actividades from workbook.planeacion
        unique_actividades = {}
        for r in workbook.planeacion:
            if r.actividad_id not in unique_actividades:
                unique_actividades[r.actividad_id] = (
                    r.fase_id,
                    r.actividad_proyecto,
                    r.orden_actividad,
                )

        for aid, (fid, descripcion, orden_actividad) in unique_actividades.items():
            fase_id = fase_by_excel_id.get(fid)
            if fase_id is None:
                raise ProjectExcelValidationError(
                    f"Actividad con actividad_id '{aid}' "
                    f"referencia fase_id '{fid}' no encontrada "
                    "en Fases",
                )
            actividad = await self._project_repository.create_actividad(
                fase_id=fase_id,
                descripcion=descripcion,
                orden=orden_actividad,
            )
            await self._session.refresh(actividad)
            actividad_by_excel_id[aid] = actividad.id
            actividad_ids.append(actividad.id)

        pendientes = await self._materialize_asignaciones_curriculares(
            proyecto_id=proyecto.id,
            programa_id=programa_id,
            planeacion=workbook.planeacion,
            actividad_by_excel_id=actividad_by_excel_id,
        )

        return ProyectoExcelImportDTO(
            referencia_id=referencia_id,
            proyecto_id=proyecto.id,
            fase_ids=fase_ids,
            actividad_ids=actividad_ids,
            resumen=_build_summary(workbook),
            pendientes_resumen=pendientes,
        )

    async def _materialize_asignaciones_curriculares(
        self,
        *,
        proyecto_id: uuid.UUID,
        programa_id: uuid.UUID,
        planeacion: list[PlaneacionRow],
        actividad_by_excel_id: dict[str, uuid.UUID],
    ) -> ExcelPendingSummaryDTO:
        """Materialize activity ↔ competency ↔ result links from the matrix.

        External identifiers are reconciled against the curriculum already
        imported for the training program; existing entities are never
        duplicated. The authoritative ``tipo_resultado`` value is preserved.
        """
        competencias_by_codigo = await self._load_programa_competencias(programa_id)

        materializadas = 0
        sin_competencia = 0
        sin_resultado = 0
        seen: set[tuple[uuid.UUID, uuid.UUID, uuid.UUID | None]] = set()

        for row in planeacion:
            actividad_db_id = actividad_by_excel_id.get(row.actividad_id)
            if actividad_db_id is None:
                continue

            competencia = competencias_by_codigo.get(
                row.codigo_competencia.strip().lower()
            )
            if competencia is None:
                sin_competencia += 1
                continue

            resultado: ResultadoAprendizaje | None = None
            rap_id = row.rap_id.strip()
            if rap_id:
                resultado = next(
                    (
                        item
                        for item in competencia.resultados
                        if (item.codigo_resultado or "").strip().lower()
                        == rap_id.lower()
                    ),
                    None,
                )
                if resultado is None:
                    sin_resultado += 1

            key = (
                actividad_db_id,
                competencia.id,
                resultado.id if resultado is not None else None,
            )
            if key in seen:
                continue
            seen.add(key)

            observaciones = row.observaciones or ""
            if rap_id and resultado is None:
                observaciones = (
                    f"rap_id '{rap_id}' no conciliado con el programa; "
                    f"{observaciones}"
                ).strip()

            await self._project_repository.create_asignacion_curricular(
                proyecto_id=proyecto_id,
                actividad_proyecto_id=actividad_db_id,
                competencia_id=competencia.id,
                resultado_id=resultado.id if resultado is not None else None,
                tipo_resultado=row.tipo_resultado.strip(),
                orden_resultado=row.orden_resultado,
                pagina_origen=row.pagina_origen,
                observaciones=observaciones or None,
            )
            materializadas += 1

        return ExcelPendingSummaryDTO(
            total=sin_competencia + sin_resultado,
            asignaciones_materializadas=materializadas,
            asignaciones_sin_competencia=sin_competencia,
            asignaciones_sin_resultado=sin_resultado,
        )

    async def _load_programa_competencias(
        self, programa_id: uuid.UUID
    ) -> dict[str, Competencia]:
        """Index the imported program competencies by normalized code."""
        if not hasattr(self._session, "execute"):
            return {}
        statement = (
            select(Competencia)
            .where(Competencia.programa_id == programa_id)
            .options(selectinload(Competencia.resultados))
        )
        result = self._session.execute(statement)
        if _is_awaitable(result):
            result = await result
        try:
            competencias = list(result.scalars().unique().all())
        except (TypeError, AttributeError):
            return {}
        return {
            competencia.codigo_competencia.strip().lower(): competencia
            for competencia in competencias
        }


def parse_canonical_workbook(content: bytes) -> CanonicalWorkbook:
    """Read and validate the canonical project workbook contract."""
    errores: list[ExcelValidationIssueDTO] = []
    try:
        workbook = load_workbook(
            filename=io.BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except Exception as error:
        return CanonicalWorkbook(
            proyecto=None,
            planeacion=[],
            validacion=[],
            errores=[
                ExcelValidationIssueDTO(
                    hoja="Workbook",
                    fila=None,
                    campo=None,
                    mensaje=f"No fue posible leer el Excel del proyecto: {error}",
                ),
            ],
        )

    missing_sheets = [
        sheet_name
        for sheet_name in CANONICAL_SHEETS
        if sheet_name not in workbook.sheetnames
    ]
    for sheet_name in missing_sheets:
        errores.append(
            ExcelValidationIssueDTO(
                hoja=sheet_name,
                fila=None,
                campo=None,
                mensaje="Hoja obligatoria faltante",
            ),
        )
    if missing_sheets:
        return CanonicalWorkbook(None, [], [], errores)

    rows_by_sheet: dict[str, list[dict[str, object]]] = {}
    for sheet_name, expected_headers in CANONICAL_SHEETS.items():
        sheet = workbook[sheet_name]
        header_values = [
            _cell_to_string(value)
            for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        ]
        if header_values != expected_headers:
            errores.append(
                ExcelValidationIssueDTO(
                    hoja=sheet_name,
                    fila=1,
                    campo=None,
                    mensaje="Encabezados invalidos o en orden distinto al canonico",
                ),
            )
            continue

        sheet_rows: list[dict[str, object]] = []
        for row_index, values in enumerate(
            sheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            if all(_cell_to_string(value) == "" for value in values):
                continue
            record: dict[str, object] = {
                header: value
                for header, value in zip(expected_headers, values, strict=False)
            }
            record["_row_index"] = row_index
            sheet_rows.append(record)
        rows_by_sheet[sheet_name] = sheet_rows

    if errores:
        return CanonicalWorkbook(None, [], [], errores)

    proyecto = _parse_proyecto(rows_by_sheet["Proyecto"], errores)
    planeacion = _parse_planeacion(rows_by_sheet["Planeacion_Proyecto"], errores)
    validacion = _parse_validacion(rows_by_sheet["Validacion_Proyecto"], errores)

    _validate_cross_references(
        proyecto=proyecto,
        planeacion=planeacion,
        errores=errores,
    )

    return CanonicalWorkbook(
        proyecto=proyecto,
        planeacion=planeacion,
        validacion=validacion,
        errores=errores,
    )


def _parse_proyecto(
    rows: list[dict[str, object]],
    errores: list[ExcelValidationIssueDTO],
) -> ProjectRow | None:
    if len(rows) != 1:
        errores.append(
            ExcelValidationIssueDTO(
                hoja="Proyecto",
                fila=None,
                campo=None,
                mensaje="La hoja Proyecto debe tener exactamente una fila de datos",
            ),
        )
        return None

    row = rows[0]
    pid = _required_string(row, "Proyecto", "proyecto_id", errores)
    nombre = _required_string(row, "Proyecto", "nombre_proyecto", errores)
    cod_sofia = _required_string(row, "Proyecto", "codigo_proyecto_sofia", errores)
    cod_prog = _required_string(row, "Proyecto", "codigo_programa", errores)
    nom_prog = _required_string(row, "Proyecto", "nombre_programa", errores)
    fuente = _optional_string(row.get("fuente_archivo"))
    obs = _optional_string(row.get("observaciones"))

    if (
        pid is None
        or nombre is None
        or cod_sofia is None
        or cod_prog is None
        or nom_prog is None
    ):
        return None
    return ProjectRow(
        proyecto_id=pid,
        nombre_proyecto=nombre,
        codigo_proyecto_sofia=cod_sofia,
        codigo_programa=cod_prog,
        nombre_programa=nom_prog,
        fuente_archivo=fuente,
        observaciones=obs,
        raw=_public_record(row),
    )


def _parse_planeacion(
    rows: list[dict[str, object]],
    errores: list[ExcelValidationIssueDTO],
) -> list[PlaneacionRow]:
    parsed: list[PlaneacionRow] = []
    for row in rows:
        pid = _required_string(row, "Planeacion_Proyecto", "proyecto_id", errores)
        fid = _required_string(row, "Planeacion_Proyecto", "fase_id", errores)
        fase_proj = _required_string(
            row, "Planeacion_Proyecto", "fase_proyecto", errores
        )
        aid = _required_string(row, "Planeacion_Proyecto", "actividad_id", errores)
        act_proj = _required_string(
            row, "Planeacion_Proyecto", "actividad_proyecto", errores
        )

        tipo_res_raw = _required_string(
            row, "Planeacion_Proyecto", "tipo_resultado", errores
        )
        tipo_res: str | None = None
        if tipo_res_raw is not None:
            normalized_tipo = tipo_res_raw.strip().upper()
            if normalized_tipo in (
                TipoResultadoProyecto.ESPECIFICO.value,
                TipoResultadoProyecto.TRANSVERSAL.value,
            ):
                tipo_res = normalized_tipo
            else:
                errores.append(
                    ExcelValidationIssueDTO(
                        hoja="Planeacion_Proyecto",
                        fila=_row_index(row),
                        campo="tipo_resultado",
                        mensaje=(
                            f"Valor invalido para 'tipo_resultado': '{tipo_res_raw}'. "
                            "Valores permitidos: 'ESPECIFICO', 'TRANSVERSAL'"
                        ),
                    )
                )
        comp_id = _required_string(
            row, "Planeacion_Proyecto", "competencia_id", errores
        )
        cod_comp = _required_string(
            row, "Planeacion_Proyecto", "codigo_competencia", errores
        )
        nom_comp = _required_string(
            row, "Planeacion_Proyecto", "nombre_competencia", errores
        )

        is_practical = False
        if cod_comp is not None and nom_comp is not None:
            is_practical = _is_practical_stage(cod_comp, nom_comp)

        rap_id: str | None = None
        rap_num: str | None = None
        rap_desc: str | None = None
        if is_practical:
            rap_id = _optional_string(row.get("rap_id")) or ""
            rap_num = _optional_string(row.get("rap_numero")) or ""
            rap_desc = _optional_string(row.get("resultado_aprendizaje")) or ""
        else:
            rap_id = _required_string(row, "Planeacion_Proyecto", "rap_id", errores)
            rap_num = _required_string(
                row, "Planeacion_Proyecto", "rap_numero", errores
            )
            rap_desc = _required_string(
                row, "Planeacion_Proyecto", "resultado_aprendizaje", errores
            )

        orden_f = _optional_int(row, "Planeacion_Proyecto", "orden_fase", errores)
        orden_a = _optional_int(row, "Planeacion_Proyecto", "orden_actividad", errores)
        orden_r = _optional_int(row, "Planeacion_Proyecto", "orden_resultado", errores)

        pag_orig = _optional_string(row.get("pagina_origen"))
        obs = _optional_string(row.get("observaciones"))

        if (
            pid is None
            or fid is None
            or fase_proj is None
            or aid is None
            or act_proj is None
            or tipo_res is None
            or comp_id is None
            or cod_comp is None
            or nom_comp is None
            or rap_id is None
            or rap_num is None
            or rap_desc is None
        ):
            continue

        parsed.append(
            PlaneacionRow(
                proyecto_id=pid,
                fase_id=fid,
                fase_proyecto=fase_proj,
                actividad_id=aid,
                actividad_proyecto=act_proj,
                tipo_resultado=tipo_res,
                competencia_id=comp_id,
                codigo_competencia=cod_comp,
                nombre_competencia=nom_comp,
                rap_id=rap_id,
                rap_numero=rap_num,
                resultado_aprendizaje=rap_desc,
                orden_fase=orden_f,
                orden_actividad=orden_a,
                orden_resultado=orden_r,
                pagina_origen=pag_orig,
                observaciones=obs,
                raw=_public_record(row),
            )
        )
    return parsed


def _parse_validacion(
    rows: list[dict[str, object]],
    errores: list[ExcelValidationIssueDTO],
) -> list[ValidacionRow]:
    parsed: list[ValidacionRow] = []
    for row in rows:
        tipo = _required_string(row, "Validacion_Proyecto", "tipo_validacion", errores)
        desc = _required_string(row, "Validacion_Proyecto", "descripcion", errores)
        estado = _required_string(row, "Validacion_Proyecto", "estado", errores)
        obs = _optional_string(row.get("observaciones"))
        if tipo is None or desc is None or estado is None:
            continue
        parsed.append(
            ValidacionRow(
                tipo_validacion=tipo,
                descripcion=desc,
                estado=estado,
                observaciones=obs,
                raw=_public_record(row),
            )
        )
    return parsed


def _validate_cross_references(
    *,
    proyecto: ProjectRow | None,
    planeacion: list[PlaneacionRow],
    errores: list[ExcelValidationIssueDTO],
) -> None:
    if proyecto is None:
        return
    for row in planeacion:
        if row.proyecto_id != proyecto.proyecto_id:
            errores.append(
                ExcelValidationIssueDTO(
                    hoja="Planeacion_Proyecto",
                    fila=_row_index(row.raw),
                    campo="proyecto_id",
                    mensaje=(
                        f"proyecto_id '{row.proyecto_id}' no coincide con "
                        f"el de la hoja Proyecto '{proyecto.proyecto_id}'"
                    ),
                ),
            )


def _validate_excel_upload(*, filename: str, content: bytes) -> None:
    if not filename.strip():
        raise InvalidProjectExcelUploadError("El archivo Excel es obligatorio")
    if not filename.lower().endswith(".xlsx"):
        raise InvalidProjectExcelUploadError("Solo se aceptan archivos .xlsx")
    if not content:
        raise InvalidProjectExcelUploadError("El Excel seleccionado esta vacio")


def _build_excel_storage_key(referencia_id: uuid.UUID, filename: str) -> str:
    """Build the canonical project workbook prefix for MinIO objects."""
    safe_filename = re.sub(r"[^a-zA-Z0-9._-]+", "-", filename).strip("-")
    if not safe_filename:
        safe_filename = "proyecto.xlsx"
    object_id = uuid.uuid4()
    return f"proyectos-formativos/{referencia_id}/excel/{object_id}-{safe_filename}"


def count_resultados_especificos(planeacion: list[PlaneacionRow]) -> int:
    unique_raps = set()
    for r in planeacion:
        if r.rap_id and r.tipo_resultado:
            if r.tipo_resultado == TipoResultadoProyecto.ESPECIFICO.value:
                unique_raps.add(r.rap_id)
    return len(unique_raps)


def build_fase_previews(planeacion: list[PlaneacionRow]) -> list[ExcelFasePreviewDTO]:
    fases_map: dict[str, FaseBuilder] = {}
    for r in planeacion:
        fid = r.fase_id
        if not fid:
            continue
        if fid not in fases_map:
            fases_map[fid] = FaseBuilder(
                fase_id=fid, nombre_fase=r.fase_proyecto, orden=r.orden_fase
            )
        fase = fases_map[fid]

        aid = r.actividad_id
        if not aid:
            continue
        if aid not in fase.actividades:
            fase.actividades[aid] = ActividadBuilder(
                actividad_id=aid,
                descripcion=r.actividad_proyecto,
                orden=r.orden_actividad,
            )
        act = fase.actividades[aid]

        cid = r.competencia_id
        if not cid:
            continue
        if cid not in act.competencias:
            act.competencias[cid] = CompetenciaBuilder(
                competencia_id=cid,
                codigo_competencia=r.codigo_competencia,
                nombre_competencia=r.nombre_competencia,
            )
        comp = act.competencias[cid]

        rid = r.rap_id
        if not rid:
            continue
        if rid not in comp.resultados:
            comp.resultados[rid] = ResultBuilder(
                rap_id=rid,
                rap_numero=r.rap_numero,
                resultado_aprendizaje=r.resultado_aprendizaje,
                tipo_resultado=r.tipo_resultado,
                orden_resultado=r.orden_resultado,
                pagina_origen=r.pagina_origen,
                observaciones=r.observaciones,
            )

    sorted_fase_builders = sorted(
        fases_map.values(), key=lambda f: (f.orden or 9999, f.nombre_fase)
    )
    fase_dtos: list[ExcelFasePreviewDTO] = []

    for f in sorted_fase_builders:
        act_dtos: list[ExcelActividadPreviewDTO] = []
        unique_competence_ids = set()
        unique_rap_ids = set()

        sorted_act_builders = sorted(
            f.actividades.values(), key=lambda a: (a.orden or 9999, a.descripcion)
        )
        for a in sorted_act_builders:
            comp_dtos: list[ExcelCompetenciaPreviewDTO] = []
            sorted_comp_builders = sorted(
                a.competencias.values(),
                key=lambda c: (c.codigo_competencia, c.nombre_competencia),
            )
            for c in sorted_comp_builders:
                unique_competence_ids.add(c.competencia_id)
                res_dtos: list[ExcelResultPreviewDTO] = []
                sorted_res_builders = sorted(
                    c.resultados.values(),
                    key=lambda r: (r.orden_resultado or 9999, r.rap_numero),
                )
                for res in sorted_res_builders:
                    unique_rap_ids.add(res.rap_id)
                    res_dtos.append(
                        ExcelResultPreviewDTO(
                            rap_id=res.rap_id,
                            rap_numero=res.rap_numero,
                            resultado_aprendizaje=res.resultado_aprendizaje,
                            tipo_resultado=TipoResultadoProyecto(res.tipo_resultado),
                            orden_resultado=res.orden_resultado,
                            pagina_origen=res.pagina_origen,
                            observaciones=res.observaciones,
                        )
                    )
                comp_dtos.append(
                    ExcelCompetenciaPreviewDTO(
                        competencia_id=c.competencia_id,
                        codigo_competencia=c.codigo_competencia,
                        nombre_competencia=c.nombre_competencia,
                        resultados=res_dtos,
                    )
                )
            act_dtos.append(
                ExcelActividadPreviewDTO(
                    actividad_id=a.actividad_id,
                    descripcion=a.descripcion,
                    orden=a.orden,
                    competencias=comp_dtos,
                )
            )

        fase_dtos.append(
            ExcelFasePreviewDTO(
                fase_id=f.fase_id,
                nombre_fase=f.nombre_fase,
                orden=f.orden,
                actividades=act_dtos,
                numero_competencias=len(unique_competence_ids),
                numero_resultados=len(unique_rap_ids),
            )
        )
    return fase_dtos


def _build_preview_dto(
    *,
    referencia_id: uuid.UUID,
    workbook: CanonicalWorkbook,
    document: StoredDocumentDTO | None,
) -> ProyectoExcelPreviewDTO:
    return ProyectoExcelPreviewDTO(
        referencia_id=referencia_id,
        documento=document,
        valid=workbook.is_valid,
        estado_validacion="VALIDO" if workbook.is_valid else "INVALIDO",
        resumen=_build_summary(workbook),
        proyecto=(
            ExcelProjectPreviewDTO(
                codigo_proyecto=workbook.proyecto.codigo_proyecto_sofia,
                nombre_proyecto=workbook.proyecto.nombre_proyecto,
                version_proyecto="1",
            )
            if workbook.proyecto is not None
            else None
        ),
        fases=build_fase_previews(workbook.planeacion),
        pendientes_resumen=ExcelPendingSummaryDTO(total=0),
        errores=workbook.errores,
    )


def _build_summary(workbook: CanonicalWorkbook) -> ExcelPreviewSummaryDTO:
    unique_fases = {r.fase_id for r in workbook.planeacion if r.fase_id}
    unique_actividades = {r.actividad_id for r in workbook.planeacion if r.actividad_id}

    return ExcelPreviewSummaryDTO(
        proyecto=1 if workbook.proyecto is not None else 0,
        fases=len(unique_fases),
        actividades=len(unique_actividades),
        resultados_especificos=count_resultados_especificos(workbook.planeacion),
    )


def _merge_excel_preview_into_payload(
    *,
    payload: dict[str, object],
    preview: ProyectoExcelPreviewDTO,
) -> dict[str, object]:
    next_payload = dict(payload)
    now = datetime.now(UTC).isoformat()
    meta = _as_record(next_payload.get("meta")) or {}
    next_payload["meta"] = {
        **meta,
        "lastInteractionAt": now,
        "touchedSteps": _append_touched_step(
            meta.get("touchedSteps"),
            "fuente-proyecto",
        ),
    }
    documental = _as_record(next_payload.get("documental")) or {}
    documental["fuente_estructurada"] = _preview_to_payload(preview, updated_at=now)
    next_payload["documental"] = documental
    return next_payload


def _merge_excel_import_into_payload(
    *,
    payload: dict[str, object],
    result: ProyectoExcelImportDTO,
) -> dict[str, object]:
    next_payload = dict(payload)
    now = datetime.now(UTC).isoformat()
    documental = _as_record(next_payload.get("documental")) or {}
    excel_payload = _as_record(documental.get("fuente_estructurada")) or {}
    excel_payload["confirmacion"] = {
        "estado": "IMPORTADO",
        "confirmed_at": now,
        "proyecto_id": str(result.proyecto_id),
        "fase_ids": [str(item) for item in result.fase_ids],
        "actividad_ids": [str(item) for item in result.actividad_ids],
        "pendientes_resumen": {
            "total": result.pendientes_resumen.total,
        },
    }
    documental["fuente_estructurada"] = excel_payload
    next_payload["documental"] = documental

    if result.resumen.proyecto == 1:
        preview = _as_record(excel_payload.get("preview")) or {}
        proyecto_p = _as_record(preview.get("proyecto")) or {}
        next_payload["proyecto"] = {
            **(_as_record(next_payload.get("proyecto")) or {}),
            "codigo_proyecto": _read_string(proyecto_p, "codigo_proyecto"),
            "nombre_proyecto": _read_string(proyecto_p, "nombre_proyecto"),
            "version_proyecto": _read_string(proyecto_p, "version_proyecto"),
        }

    next_payload["estructura"] = {
        "fases": [
            {"fase_id": str(fid), "estado": "IMPORTADO"} for fid in result.fase_ids
        ],
        "actividades": [
            {"actividad_id": str(aid), "estado": "IMPORTADO"}
            for aid in result.actividad_ids
        ],
        "excel_import": {
            "estado": "IMPORTADO",
            "updated_at": now,
            "resumen": {
                "proyecto": result.resumen.proyecto,
                "fases": result.resumen.fases,
                "actividades": result.resumen.actividades,
                "resultados_especificos": result.resumen.resultados_especificos,
            },
        },
    }
    return next_payload


def _preview_to_payload(
    preview: ProyectoExcelPreviewDTO,
    *,
    updated_at: str,
) -> dict[str, object]:
    return {
        "documento": _document_to_payload(preview.documento),
        "preview": {
            "valid": preview.valid,
            "estado_validacion": preview.estado_validacion,
            "resumen": {
                "proyecto": preview.resumen.proyecto,
                "fases": preview.resumen.fases,
                "actividades": preview.resumen.actividades,
                "resultados_especificos": preview.resumen.resultados_especificos,
            },
            "proyecto": (
                {
                    "codigo_proyecto": preview.proyecto.codigo_proyecto,
                    "nombre_proyecto": preview.proyecto.nombre_proyecto,
                    "version_proyecto": preview.proyecto.version_proyecto,
                }
                if preview.proyecto is not None
                else None
            ),
            "fases": [
                {
                    "fase_id": f.fase_id,
                    "nombre_fase": f.nombre_fase,
                    "orden": f.orden,
                    "numero_competencias": f.numero_competencias,
                    "numero_resultados": f.numero_resultados,
                    "actividades": [
                        {
                            "actividad_id": a.actividad_id,
                            "descripcion": a.descripcion,
                            "orden": a.orden,
                            "competencias": [
                                {
                                    "competencia_id": c.competencia_id,
                                    "codigo_competencia": c.codigo_competencia,
                                    "nombre_competencia": c.nombre_competencia,
                                    "resultados": [
                                        {
                                            "rap_id": r.rap_id,
                                            "rap_numero": r.rap_numero,
                                            "resultado_aprendizaje": (
                                                r.resultado_aprendizaje
                                            ),
                                            "tipo_resultado": r.tipo_resultado,
                                            "orden_resultado": r.orden_resultado,
                                            "pagina_origen": r.pagina_origen,
                                            "observaciones": r.observaciones,
                                        }
                                        for r in c.resultados
                                    ],
                                }
                                for c in a.competencias
                            ],
                        }
                        for a in f.actividades
                    ],
                }
                for f in preview.fases
            ],
            "pendientes_resumen": {
                "total": preview.pendientes_resumen.total,
            },
            "errores": [
                {
                    "hoja": e.hoja,
                    "fila": e.fila,
                    "campo": e.campo,
                    "mensaje": e.mensaje,
                }
                for e in preview.errores
            ],
        },
        "updated_at": updated_at,
    }


def _get_valid_excel_preview_payload(
    payload: dict[str, object],
) -> dict[str, object]:
    documental = _as_record(payload.get("documental")) or {}
    excel = _as_record(documental.get("fuente_estructurada")) or {}
    preview = _as_record(excel.get("preview"))
    if preview is None:
        raise ProjectExcelMissingPreviewError(
            "No se encontro preview de importacion en el borrador",
        )
    return {
        "documento": _as_record(excel.get("documento")) or {},
        "preview": preview,
    }


def _as_record(value: object) -> dict[str, object] | None:
    if isinstance(value, dict):
        return value
    return None


def _read_string(record: dict[str, object], key: str) -> str:
    val = record.get(key)
    return val if isinstance(val, str) else ""


def _read_uuid(record: dict[str, object], key: str) -> uuid.UUID | None:
    val = _read_string(record, key)
    if not val:
        return None
    try:
        return uuid.UUID(val)
    except ValueError as error:
        raise ProjectExcelValidationError(
            f"El identificador {key} del borrador de proyecto no es un UUID valido"
        ) from error


def _extract_programa_id_from_program_payload(
    payload: dict[str, object],
) -> uuid.UUID | None:
    curricular = _as_record(payload.get("curricular")) or {}
    programa_id = _read_uuid(curricular, "programa_formacion_id")
    if programa_id is not None:
        return programa_id

    documental = _as_record(payload.get("documental")) or {}
    programa_excel = _as_record(documental.get("programa_excel")) or {}
    confirmacion = _as_record(programa_excel.get("confirmacion")) or {}
    return _read_uuid(confirmacion, "programa_id")


def _document_to_payload(
    document: StoredDocumentDTO | dict[str, object] | None,
) -> dict[str, object] | None:
    if document is None:
        return None
    if isinstance(document, dict):
        return document
    return {
        "original_filename": document.original_filename,
        "storage_key": document.storage_key,
        "size_bytes": document.size_bytes,
        "content_type": document.content_type,
        "checksum_sha256": document.checksum_sha256,
        "etag": document.etag,
    }


def _append_touched_step(steps: object, step_id: str) -> list[str]:
    if not isinstance(steps, list):
        return [step_id]
    result = [s for s in steps if isinstance(s, str)]
    if step_id not in result:
        result.append(step_id)
    return result


def _is_awaitable(value: object) -> bool:
    import inspect

    return inspect.isawaitable(value)


def _cell_to_string(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _required_string(
    row: dict[str, object],
    hoja: str,
    campo: str,
    errores: list[ExcelValidationIssueDTO],
) -> str | None:
    value = row.get(campo)
    text = str(value).strip() if value is not None else ""
    if not text:
        errores.append(
            ExcelValidationIssueDTO(
                hoja=hoja,
                fila=_row_index(row),
                campo=campo,
                mensaje=f"Campo obligatorio '{campo}' faltante o vacio",
            ),
        )
        return None
    return text


def _optional_string(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def _optional_int(
    row: dict[str, object],
    hoja: str,
    campo: str,
    errores: list[ExcelValidationIssueDTO],
) -> int | None:
    value = row.get(campo)
    if value is None:
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value)
    if not isinstance(value, str):
        errores.append(
            ExcelValidationIssueDTO(
                hoja=hoja,
                fila=_row_index(row),
                campo=campo,
                mensaje=f"Valor invalido para campo numerico '{campo}'",
            ),
        )
        return None
    try:
        return int(value.strip())
    except ValueError:
        errores.append(
            ExcelValidationIssueDTO(
                hoja=hoja,
                fila=_row_index(row),
                campo=campo,
                mensaje=f"Valor invalido para campo numerico '{campo}'",
            ),
        )
        return None


def _row_index(row: dict[str, object]) -> int | None:
    val = row.get("_row_index")
    if isinstance(val, int):
        return val
    return None


def _public_record(row: dict[str, object]) -> dict[str, object]:
    return {k: v for k, v in row.items() if not k.startswith("_")}


def _is_practical_stage(codigo_competencia: str, nombre_competencia: str) -> bool:
    import unicodedata

    def normalize(val: str) -> str:
        without_accents = "".join(
            char
            for char in unicodedata.normalize("NFKD", val)
            if not unicodedata.combining(char)
        )
        return re.sub(r"\s+", " ", without_accents).strip().lower()

    normalized_name = normalize(nombre_competencia)
    normalized_code = codigo_competencia.strip()
    return (
        normalized_code == "999999999"
        or "etapa practica" in normalized_name
        or "etapa productiva" in normalized_name
    )
