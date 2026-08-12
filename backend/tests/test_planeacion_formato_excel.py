"""Tests for the official GPFI-F-134 V05 workbook generator."""

from __future__ import annotations

import hashlib
import io
from datetime import date
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from src.application.services.planeacion_formato_excel import (
    EXCEL_CONTENT_TYPE,
    TEMPLATE_PATH,
    FormatoPlaneacionMetadata,
    FormatoPlaneacionRow,
    PlaneacionFormatoExcelService,
    PlaneacionFormatoValidationError,
)


def _metadata() -> FormatoPlaneacionMetadata:
    return FormatoPlaneacionMetadata(
        fecha_elaboracion=date(2026, 7, 29),
        nombre_programa="Analisis y desarrollo de software",
        modalidad_formacion="Presencial",
        codigo_programa="228118",
        version_programa="1",
        nombre_proyecto="Sistema de informacion",
        codigo_proyecto="PR-001",
        equipo_gestion_curricular=("Ana Instructor", "Luis Instructor"),
        regional="Distrito Capital",
        centro_formacion="Centro de Servicios Financieros",
        clasificacion_informacion="PUBLICA_CLASIFICADA",
    )


def _row(index: int = 1) -> FormatoPlaneacionRow:
    return FormatoPlaneacionRow(
        fase=f"Fase {index}",
        actividad_proyecto=f"Actividad {index}",
        competencia="220501046\nDesarrollar software",
        resultado=f"RAP-{index:02d}\nResultado de aprendizaje {index}",
        saberes=("Saber uno", "Saber dos"),
        procesos=("Proceso uno",),
        criterios=("Criterio uno",),
        actividades_aprendizaje="Resolver un reto aplicado",
        horas_trabajo_directo=8,
        horas_trabajo_independiente=4,
        descripcion_evidencia="Producto funcional",
        estrategias_didacticas="Aprendizaje basado en proyectos",
        ambiente=("Aula TIC", "Laboratorio"),
        materiales_formacion="Computador y repositorio",
        instructores="Ana Instructor",
        observaciones="Sin observaciones",
    )


def test_template_is_packaged_and_has_official_sheets() -> None:
    assert TEMPLATE_PATH.is_file()
    workbook = load_workbook(TEMPLATE_PATH, read_only=True)
    assert workbook.sheetnames == ["Instrucciones", "FASE"]


def test_generator_preserves_package_and_fills_official_mapping() -> None:
    result = PlaneacionFormatoExcelService().generar(
        metadata=_metadata(),
        rows=[_row()],
    )
    workbook = load_workbook(io.BytesIO(result.content))
    sheet = workbook["FASE"]

    assert workbook.sheetnames == ["Instrucciones", "FASE"]
    assert sheet["E10"].value == "Analisis y desarrollo de software"
    assert sheet["E11"].value == "Presencial"
    assert sheet["E12"].value == "228118 / 1"
    assert sheet["A18"].value == "Fase 1"
    assert sheet["E18"].value == "- Saber uno\n- Saber dos"
    assert sheet["F18"].value == "- Proceso uno"
    assert sheet["I18"].value == 8
    assert sheet["J18"].value == 4
    assert sheet.print_area == "'FASE'!$A$1:$P$18"
    assert len(sheet.merged_cells.ranges) == 37
    assert hashlib.sha256(result.content).hexdigest() == result.checksum_sha256
    assert EXCEL_CONTENT_TYPE.endswith("spreadsheetml.sheet")

    with ZipFile(TEMPLATE_PATH) as template_zip:
        with ZipFile(io.BytesIO(result.content)) as result_zip:
            assert (
                result_zip.read("xl/worksheets/sheet1.xml")
                == template_zip.read("xl/worksheets/sheet1.xml")
            )
            for prefix in (
                "xl/drawings/",
                "xl/media/",
                "xl/printerSettings/",
            ):
                assert any(
                    name.startswith(prefix) for name in result_zip.namelist()
                )
            drawing = result_zip.read("xl/drawings/drawing2.xml")
            assert b">X<" in drawing


def test_generator_inserts_additional_styled_rows() -> None:
    result = PlaneacionFormatoExcelService().generar(
        metadata=_metadata(),
        rows=[_row(index) for index in range(1, 14)],
    )
    workbook = load_workbook(io.BytesIO(result.content))
    sheet = workbook["FASE"]

    assert sheet["A30"].value == "Fase 13"
    assert sheet["P30"].value == "Sin observaciones"
    assert sheet["A30"].style_id == sheet["A18"].style_id
    assert sheet.print_area == "'FASE'!$A$1:$P$30"


def test_generator_preserves_ignorable_namespace_prefixes() -> None:
    result = PlaneacionFormatoExcelService().generar(
        metadata=_metadata(),
        rows=[_row()],
    )

    with ZipFile(io.BytesIO(result.content)) as package:
        workbook_xml = package.read("xl/workbook.xml")

    assert b'mc:Ignorable="x15 xr xr6 xr10 xr2"' in workbook_xml
    for prefix in (
        b"xmlns:x15=",
        b"xmlns:xr=",
        b"xmlns:xr6=",
        b"xmlns:xr10=",
        b"xmlns:xr2=",
    ):
        assert prefix in workbook_xml


def test_generator_rejects_negative_hours() -> None:
    invalid = _row()
    invalid = FormatoPlaneacionRow(
        **{
            **invalid.__dict__,
            "horas_trabajo_directo": -1,
        }
    )
    with pytest.raises(PlaneacionFormatoValidationError):
        PlaneacionFormatoExcelService().generar(
            metadata=_metadata(),
            rows=[invalid],
        )


def _block_row(
    *,
    resultado: str,
    horas_directo: float | None,
    horas_independiente: float | None,
) -> FormatoPlaneacionRow:
    base = _row()
    return FormatoPlaneacionRow(
        **{
            **base.__dict__,
            "resultado": resultado,
            "horas_trabajo_directo": horas_directo,
            "horas_trabajo_independiente": horas_independiente,
        }
    )


def test_generator_multirap_block_writes_hours_once() -> None:
    """A 20h learning activity split across 3 RAPs must not sum 60h.

    The hours belong to the integrated activity, so only the first row of
    the block carries them; the remaining RAP rows stay blank.
    """
    rows = [
        _block_row(
            resultado="RAP-01\nResultado tecnico uno",
            horas_directo=12,
            horas_independiente=8,
        ),
        _block_row(
            resultado="RAP-02\nResultado tecnico dos",
            horas_directo=None,
            horas_independiente=None,
        ),
        _block_row(
            resultado="RAP-03\nResultado transversal",
            horas_directo=None,
            horas_independiente=None,
        ),
    ]
    result = PlaneacionFormatoExcelService().generar(
        metadata=_metadata(),
        rows=rows,
    )
    workbook = load_workbook(io.BytesIO(result.content))
    sheet = workbook["FASE"]

    first_row, second_row, third_row = 18, 19, 20
    assert sheet.cell(first_row, 4).value == "RAP-01\nResultado tecnico uno"
    assert sheet.cell(second_row, 4).value == "RAP-02\nResultado tecnico dos"
    assert sheet.cell(third_row, 4).value == "RAP-03\nResultado transversal"

    assert sheet.cell(first_row, 9).value == 12
    assert sheet.cell(first_row, 10).value == 8
    assert sheet.cell(second_row, 9).value is None
    assert sheet.cell(second_row, 10).value is None
    assert sheet.cell(third_row, 9).value is None
    assert sheet.cell(third_row, 10).value is None

    total_directo = sum(
        value
        for value in (
            sheet.cell(first_row, 9).value,
            sheet.cell(second_row, 9).value,
            sheet.cell(third_row, 9).value,
        )
        if isinstance(value, (int, float))
    )
    total_independiente = sum(
        value
        for value in (
            sheet.cell(first_row, 10).value,
            sheet.cell(second_row, 10).value,
            sheet.cell(third_row, 10).value,
        )
        if isinstance(value, (int, float))
    )
    assert total_directo == 12
    assert total_independiente == 8
    assert result.filas_generadas == 3


def test_generator_multirap_block_keeps_shared_block_fields() -> None:
    """Phase, activity and learning activity repeat across block rows."""
    rows = [
        _block_row(
            resultado="RAP-01\nResultado tecnico",
            horas_directo=4,
            horas_independiente=2,
        ),
        _block_row(
            resultado="RAP-02\nResultado transversal",
            horas_directo=None,
            horas_independiente=None,
        ),
    ]
    result = PlaneacionFormatoExcelService().generar(
        metadata=_metadata(),
        rows=rows,
    )
    workbook = load_workbook(io.BytesIO(result.content))
    sheet = workbook["FASE"]

    assert sheet.cell(18, 1).value == sheet.cell(19, 1).value == "Fase 1"
    assert sheet.cell(18, 2).value == sheet.cell(19, 2).value == "Actividad 1"
    assert (
        sheet.cell(18, 8).value
        == sheet.cell(19, 8).value
        == "Resolver un reto aplicado"
    )
