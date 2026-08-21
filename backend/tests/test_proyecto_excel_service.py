"""Tests for project Excel import service."""

from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime
from unittest.mock import MagicMock

import openpyxl
import pytest
from openpyxl import Workbook

from src.application.dto.programa_documentos import StoredDocumentDTO
from src.application.services.proyecto_excel import (
    InvalidProjectExcelUploadError,
    PlaneacionRow,
    ProjectExcelDraftMissingError,
    ProjectExcelMissingPreviewError,
    ProjectExcelValidationError,
    ProyectoExcelImportService,
    count_resultados_especificos,
    parse_canonical_workbook,
)
from src.domain.shared.enums import EstadoBloque
from src.infrastructure.db.models.curriculum import (
    Competencia,
    ProgramaFormacion,
    ResultadoAprendizaje,
)

pytestmark = pytest.mark.anyio


def _build_competencia(
    *,
    codigo_competencia: str,
    nombre: str,
    rap_codigos: list[str],
) -> Competencia:
    competencia = Competencia(
        codigo_competencia=codigo_competencia,
        nombre_competencia=nombre,
    )
    competencia.id = uuid.uuid4()
    competencia.resultados = []
    for rap_codigo in rap_codigos:
        resultado = ResultadoAprendizaje(
            codigo_resultado=rap_codigo,
            descripcion=f"Resultado {rap_codigo}",
        )
        resultado.id = uuid.uuid4()
        resultado.competencia_id = competencia.id
        competencia.resultados.append(resultado)
    return competencia


class MockDraftRepository:
    def __init__(self, draft=None, program_draft=None):
        self.draft = draft
        self.program_draft = program_draft
        self.saved = False

    async def get_by_block_reference(self, tipo_bloque, referencia_id):
        if getattr(tipo_bloque, "value", tipo_bloque) == "PROGRAMA":
            return self.program_draft
        return self.draft

    async def save(self, draft):
        self.saved = True
        return draft


class MockAuditRepository:
    def __init__(self):
        self.events = []

    async def add_event(self, entidad, entidad_id, accion, detalle=None):
        self.events.append({"entidad": entidad, "accion": accion, "detalle": detalle})


class MockSession:
    def __init__(self, competencias=None):
        self.committed = False
        self.refreshed = False
        self._competencias = competencias or []
        self._programa = MagicMock(spec=ProgramaFormacion)

    async def commit(self):
        self.committed = True

    async def refresh(self, instance):
        self.refreshed = True

    async def execute(self, statement):
        entity = None
        try:
            entity = statement.column_descriptions[0].get("entity")
        except Exception:
            entity = None
        result = MagicMock()
        if entity is Competencia:
            scalars = MagicMock()
            scalars.unique.return_value.all.return_value = self._competencias
            result.scalars.return_value = scalars
        else:
            result.scalar_one_or_none.return_value = self._programa
        return result


class MockStorageService:
    def __init__(self, document=None, content=b""):
        self.document = document or StoredDocumentDTO(
            original_filename="proyecto.xlsx",
            storage_key="proyectos-formativos/test-id/excel/test.xlsx",
            size_bytes=1024,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            checksum_sha256="abc123",
            etag="etag-123",
        )
        self._content = content

    async def save_excel(self, *, key, content, content_type, original_filename):
        self._content = content
        return self.document

    async def read_excel(self, *, key):
        if not self._content:
            raise FileNotFoundError("Excel no encontrado")
        return self._content


class MockProjectRepository:
    def __init__(
        self,
        existing_codigo_proyecto=None,
        existing_nombre_proyecto=None,
    ):
        self.created_proyectos = []
        self.created_fases = []
        self.created_actividades = []
        self.created_asignaciones = []
        self.created_programa_ids = []
        self.existing_codigo_proyecto = existing_codigo_proyecto
        self.existing_nombre_proyecto = existing_nombre_proyecto

    async def create_proyecto(
        self,
        *,
        programa_id,
        codigo_proyecto,
        nombre_proyecto,
        version_proyecto,
    ):
        p = MagicMock(id=uuid.uuid4())
        p.codigo_proyecto = codigo_proyecto
        self.created_programa_ids.append(programa_id)
        self.created_proyectos.append(p)
        return p

    async def create_fase(self, *, proyecto_id, nombre_fase, orden):
        f = MagicMock(id=uuid.uuid4())
        f.nombre_fase = nombre_fase
        self.created_fases.append(f)
        return f

    async def create_actividad(self, *, fase_id, descripcion, orden):
        a = MagicMock(id=uuid.uuid4())
        a.descripcion = descripcion
        self.created_actividades.append(a)
        return a

    async def create_asignacion_curricular(
        self,
        *,
        proyecto_id,
        actividad_proyecto_id,
        competencia_id,
        resultado_id,
        tipo_resultado,
        orden_resultado,
        pagina_origen,
        observaciones,
    ):
        asignacion = MagicMock(id=uuid.uuid4())
        asignacion.proyecto_id = proyecto_id
        asignacion.actividad_proyecto_id = actividad_proyecto_id
        asignacion.competencia_id = competencia_id
        asignacion.resultado_id = resultado_id
        asignacion.tipo_resultado = tipo_resultado
        asignacion.orden_resultado = orden_resultado
        self.created_asignaciones.append(asignacion)
        return asignacion

    async def proyecto_exists_by_code_name(
        self,
        *,
        programa_id,
        codigo_proyecto,
        nombre_proyecto,
    ):
        return (
            self.existing_codigo_proyecto is not None
            and self.existing_nombre_proyecto is not None
            and self.existing_codigo_proyecto.lower() == codigo_proyecto.lower()
            and self.existing_nombre_proyecto.lower() == nombre_proyecto.lower()
        )


def create_valid_workbook():
    wb = Workbook()

    proyecto_sheet = wb.active
    proyecto_sheet.title = "Proyecto"
    proyecto_sheet.append(
        [
            "proyecto_id",
            "nombre_proyecto",
            "codigo_proyecto_sofia",
            "codigo_programa",
            "nombre_programa",
            "fuente_archivo",
            "observaciones",
        ],
    )
    proyecto_sheet.append(
        [
            "PROJ-01",
            "Proyecto de prueba",
            "PR-001",
            "prog-ref",
            "Programa de prueba",
            "fuente.pdf",
            "observaciones proyecto",
        ]
    )

    planeacion_sheet = wb.create_sheet("Planeacion_Proyecto")
    planeacion_sheet.append(
        [
            "proyecto_id",
            "fase_id",
            "fase_proyecto",
            "actividad_id",
            "actividad_proyecto",
            "tipo_resultado",
            "competencia_id",
            "codigo_competencia",
            "nombre_competencia",
            "rap_id",
            "rap_numero",
            "resultado_aprendizaje",
            "orden_fase",
            "orden_actividad",
            "orden_resultado",
            "pagina_origen",
            "observaciones",
        ],
    )
    planeacion_sheet.append(
        [
            "PROJ-01",
            "F1",
            "Fase 1",
            "A1",
            "Actividad 1.1",
            "ESPECIFICO",
            "C1",
            "240201050",
            "Competencia de prueba 1",
            "R1",
            "1",
            "Resultado de aprendizaje 1",
            1,
            1,
            1,
            "12",
            "",
        ]
    )
    planeacion_sheet.append(
        [
            "PROJ-01",
            "F1",
            "Fase 1",
            "A2",
            "Actividad 1.2",
            "ESPECIFICO",
            "C1",
            "240201050",
            "Competencia de prueba 1",
            "R2",
            "2",
            "Resultado de aprendizaje 2",
            1,
            2,
            2,
            "13",
            "",
        ]
    )
    planeacion_sheet.append(
        [
            "PROJ-01",
            "F2",
            "Fase 2",
            "A3",
            "Actividad 2.1",
            "TRANSVERSAL",
            "C2",
            "240201051",
            "Competencia de prueba 2",
            "R3",
            "1",
            "Resultado de aprendizaje 3",
            2,
            1,
            1,
            "15",
            "",
        ]
    )

    validacion_sheet = wb.create_sheet("Validacion_Proyecto")
    validacion_sheet.append(
        [
            "tipo_validacion",
            "descripcion",
            "estado",
            "observaciones",
        ]
    )
    validacion_sheet.append(
        [
            "CURRICULAR",
            "Validacion basica del diseno",
            "OK",
            "Ninguna",
        ]
    )

    return wb


@pytest.fixture
def draft_factory():
    def _build():
        return MagicMock(
            id=uuid.uuid4(),
            payload_json={
                "meta": {
                    "referenciaId": str(uuid.uuid4()),
                    "programaId": "00000000-0000-0000-0000-000000000000",
                    "touchedSteps": ["fuente-proyecto"],
                    "lastInteractionAt": datetime.now(UTC).isoformat(),
                },
                "documental": {},
            },
            paso_actual="fuente-proyecto",
            estado_borrador=EstadoBloque.BORRADOR,
        )

    return _build


@pytest.fixture
def default_competencias():
    return [
        _build_competencia(
            codigo_competencia="240201050",
            nombre="Competencia de prueba 1",
            rap_codigos=["R1", "R2"],
        ),
        _build_competencia(
            codigo_competencia="240201051",
            nombre="Competencia de prueba 2",
            rap_codigos=["R3"],
        ),
    ]


@pytest.fixture
def service(draft_factory, default_competencias):
    session = MockSession(competencias=default_competencias)
    draft_repository = MockDraftRepository(draft_factory())
    audit_repository = MockAuditRepository()
    storage_service = MockStorageService()
    project_repository = MockProjectRepository()

    return ProyectoExcelImportService(
        session=session,
        draft_repository=draft_repository,
        audit_repository=audit_repository,
        project_repository=project_repository,
        storage_service=storage_service,
    )


class TestPreviewProjectExcel:
    async def test_preview_accepts_valid_workbook(self, service):
        wb = create_valid_workbook()
        content = io.BytesIO()
        wb.save(content)
        content.seek(0)

        result = await service.preview_project_excel(
            referencia_id=uuid.uuid4(),
            filename="proyecto.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            content=content.read(),
        )

        assert result.valid is True
        assert result.estado_validacion == "VALIDO"
        assert result.proyecto is not None
        assert result.proyecto.codigo_proyecto == "PR-001"
        assert len(result.fases) == 2

    async def test_preview_rejects_missing_sheets(self, service):
        wb = Workbook()
        proyecto_sheet = wb.active
        proyecto_sheet.title = "Proyecto"
        proyecto_sheet.append(
            [
                "proyecto_id",
                "nombre_proyecto",
                "codigo_proyecto_sofia",
                "codigo_programa",
                "nombre_programa",
                "fuente_archivo",
                "observaciones",
            ]
        )
        proyecto_sheet.append(
            [
                "PROJ-01",
                "Proyecto de prueba",
                "PR-001",
                "prog-99",
                "Prog Name",
                "test.pdf",
                "",
            ]
        )
        content = io.BytesIO()
        wb.save(content)
        content.seek(0)

        result = await service.preview_project_excel(
            referencia_id=uuid.uuid4(),
            filename="proyecto.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            content=content.read(),
        )

        assert result.valid is False
        assert result.estado_validacion == "INVALIDO"
        assert len(result.errores) > 0

    async def test_preview_rejects_non_xlsx(self, service):
        with pytest.raises(
            InvalidProjectExcelUploadError,
            match="Solo se aceptan archivos .xlsx",
        ):
            await service.preview_project_excel(
                referencia_id=uuid.uuid4(),
                filename="proyecto.pdf",
                content_type="application/pdf",
                content=b"fake pdf",
            )

    async def test_preview_rejects_empty_content(self, service):
        with pytest.raises(InvalidProjectExcelUploadError, match="esta vacio"):
            await service.preview_project_excel(
                referencia_id=uuid.uuid4(),
                filename="proyecto.xlsx",
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                content=b"",
            )

    async def test_preview_rejects_missing_draft(self):
        session = MockSession()
        draft_repository = MockDraftRepository(draft=None)
        audit_repository = MockAuditRepository()
        storage_service = MockStorageService()
        project_repository = MockProjectRepository()

        service = ProyectoExcelImportService(
            session=session,
            draft_repository=draft_repository,
            audit_repository=audit_repository,
            project_repository=project_repository,
            storage_service=storage_service,
        )

        wb = create_valid_workbook()
        content = io.BytesIO()
        wb.save(content)
        content.seek(0)

        with pytest.raises(ProjectExcelDraftMissingError):
            await service.preview_project_excel(
                referencia_id=uuid.uuid4(),
                filename="proyecto.xlsx",
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                content=content.read(),
            )

    async def test_preview_rejects_existing_project_with_same_code_and_name(
        self,
        draft_factory,
    ):
        session = MockSession()
        draft_repository = MockDraftRepository(draft_factory())
        audit_repository = MockAuditRepository()
        storage_service = MockStorageService()
        project_repository = MockProjectRepository(
            existing_codigo_proyecto="PR-001",
            existing_nombre_proyecto="Proyecto de prueba",
        )
        service = ProyectoExcelImportService(
            session=session,
            draft_repository=draft_repository,
            audit_repository=audit_repository,
            project_repository=project_repository,
            storage_service=storage_service,
        )
        wb = create_valid_workbook()
        content = io.BytesIO()
        wb.save(content)
        content.seek(0)

        with pytest.raises(ProjectExcelValidationError, match="ya fue cargado"):
            await service.preview_project_excel(
                referencia_id=uuid.uuid4(),
                filename="proyecto.xlsx",
                content_type=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
                content=content.read(),
            )

        assert project_repository.created_proyectos == []

    async def test_preview_allows_new_project_when_code_and_name_differ(
        self,
        draft_factory,
    ):
        session = MockSession()
        draft_repository = MockDraftRepository(draft_factory())
        audit_repository = MockAuditRepository()
        storage_service = MockStorageService()
        project_repository = MockProjectRepository(
            existing_codigo_proyecto="PR-999",
            existing_nombre_proyecto="Proyecto anterior",
        )
        service = ProyectoExcelImportService(
            session=session,
            draft_repository=draft_repository,
            audit_repository=audit_repository,
            project_repository=project_repository,
            storage_service=storage_service,
        )
        wb = create_valid_workbook()
        content = io.BytesIO()
        wb.save(content)
        content.seek(0)

        result = await service.preview_project_excel(
            referencia_id=uuid.uuid4(),
            filename="proyecto.xlsx",
            content_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            content=content.read(),
        )

        assert result.valid is True
        assert result.proyecto is not None
        assert result.proyecto.codigo_proyecto == "PR-001"

    async def test_preview_persists_in_draft(self, service):
        wb = create_valid_workbook()
        content = io.BytesIO()
        wb.save(content)
        content.seek(0)

        await service.preview_project_excel(
            referencia_id=uuid.uuid4(),
            filename="proyecto.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            content=content.read(),
        )

        assert service._draft_repository.saved is True

    async def test_preview_sets_paso_actual(self, service):
        wb = create_valid_workbook()
        content = io.BytesIO()
        wb.save(content)
        content.seek(0)

        await service.preview_project_excel(
            referencia_id=uuid.uuid4(),
            filename="proyecto.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            content=content.read(),
        )

        assert service._draft_repository.draft.paso_actual == "fuente-proyecto"

    async def test_preview_triggers_audit_event(self, service):
        wb = create_valid_workbook()
        content = io.BytesIO()
        wb.save(content)
        content.seek(0)

        await service.preview_project_excel(
            referencia_id=uuid.uuid4(),
            filename="proyecto.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            content=content.read(),
        )

        assert len(service._audit_repository.events) == 1
        event = service._audit_repository.events[0]
        assert event["accion"] == "EXCEL_PROYECTO_PREVISUALIZADO"


class TestConfirmProjectExcelImport:
    async def test_confirm_materializes_proyecto(self, service):
        wb = create_valid_workbook()
        content = io.BytesIO()
        wb.save(content)
        content.seek(0)

        await service.preview_project_excel(
            referencia_id=uuid.uuid4(),
            filename="proyecto.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            content=content.read(),
        )

        result = await service.confirm_project_excel_import(
            referencia_id=uuid.uuid4(),
        )

        assert result.proyecto_id is not None
        assert result.resumen.proyecto == 1
        assert result.resumen.fases == 2
        assert result.resumen.actividades == 3

    async def test_confirm_materializes_fases(self, service):
        wb = create_valid_workbook()
        content = io.BytesIO()
        wb.save(content)
        content.seek(0)

        await service.preview_project_excel(
            referencia_id=uuid.uuid4(),
            filename="proyecto.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            content=content.read(),
        )

        result = await service.confirm_project_excel_import(
            referencia_id=uuid.uuid4(),
        )

        assert len(result.fase_ids) == 2

    async def test_confirm_materializes_actividades(self, service):
        wb = create_valid_workbook()
        content = io.BytesIO()
        wb.save(content)
        content.seek(0)

        await service.preview_project_excel(
            referencia_id=uuid.uuid4(),
            filename="proyecto.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            content=content.read(),
        )

        result = await service.confirm_project_excel_import(
            referencia_id=uuid.uuid4(),
        )

        assert len(result.actividad_ids) == 3

    async def test_confirm_resolves_programa_id_from_program_draft(self):
        programa_id = uuid.uuid4()
        programa_referencia_id = uuid.uuid4()
        proyecto_referencia_id = uuid.uuid4()
        proyecto_draft = MagicMock(
            id=uuid.uuid4(),
            payload_json={
                "meta": {
                    "referenciaId": str(proyecto_referencia_id),
                    "programaReferenciaId": str(programa_referencia_id),
                    "touchedSteps": ["fuente-proyecto"],
                    "lastInteractionAt": datetime.now(UTC).isoformat(),
                },
                "documental": {},
            },
            paso_actual="fuente-proyecto",
            estado_borrador=EstadoBloque.BORRADOR,
        )
        programa_draft = MagicMock(
            payload_json={
                "meta": {"referenciaId": str(programa_referencia_id)},
                "curricular": {"programa_formacion_id": str(programa_id)},
            },
        )
        session = MockSession()
        draft_repository = MockDraftRepository(proyecto_draft, programa_draft)
        audit_repository = MockAuditRepository()
        storage_service = MockStorageService()
        project_repository = MockProjectRepository()
        service = ProyectoExcelImportService(
            session=session,
            draft_repository=draft_repository,
            audit_repository=audit_repository,
            project_repository=project_repository,
            storage_service=storage_service,
        )
        wb = create_valid_workbook()
        content = io.BytesIO()
        wb.save(content)
        content.seek(0)

        await service.preview_project_excel(
            referencia_id=proyecto_referencia_id,
            filename="proyecto.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            content=content.read(),
        )

        await service.confirm_project_excel_import(
            referencia_id=proyecto_referencia_id,
        )

        assert project_repository.created_programa_ids == [programa_id]
        assert proyecto_draft.payload_json["meta"]["programaId"] == str(programa_id)

    async def test_confirm_rejects_missing_preview(self):
        session = MockSession()
        draft_repository = MockDraftRepository(
            MagicMock(
                id=uuid.uuid4(),
                payload_json={
                    "meta": {
                        "programaId": "00000000-0000-0000-0000-000000000000",
                    },
                    "documental": {},
                },
                paso_actual="fuente-proyecto",
                estado_borrador=EstadoBloque.BORRADOR,
            )
        )
        audit_repository = MockAuditRepository()
        storage_service = MockStorageService()
        project_repository = MockProjectRepository()

        service = ProyectoExcelImportService(
            session=session,
            draft_repository=draft_repository,
            audit_repository=audit_repository,
            project_repository=project_repository,
            storage_service=storage_service,
        )

        with pytest.raises(ProjectExcelMissingPreviewError):
            await service.confirm_project_excel_import(
                referencia_id=uuid.uuid4(),
            )

    async def test_confirm_commits_session(self, service):
        wb = create_valid_workbook()
        content = io.BytesIO()
        wb.save(content)
        content.seek(0)

        await service.preview_project_excel(
            referencia_id=uuid.uuid4(),
            filename="proyecto.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            content=content.read(),
        )

        await service.confirm_project_excel_import(
            referencia_id=uuid.uuid4(),
        )

        assert service._session.committed is True

    async def test_confirm_stores_excel_in_minio(self):
        storage = MockStorageService()
        session = MockSession()
        draft_repository = MockDraftRepository(
            MagicMock(
                id=uuid.uuid4(),
                payload_json={
                    "meta": {
                        "programaId": "00000000-0000-0000-0000-000000000000",
                    },
                    "documental": {},
                },
                paso_actual="fuente-proyecto",
                estado_borrador=EstadoBloque.BORRADOR,
            )
        )
        audit = MockAuditRepository()
        project_repository = MockProjectRepository()

        service = ProyectoExcelImportService(
            session=session,
            draft_repository=draft_repository,
            audit_repository=audit,
            project_repository=project_repository,
            storage_service=storage,
        )

        wb = create_valid_workbook()
        content = io.BytesIO()
        wb.save(content)
        content.seek(0)

        await service.preview_project_excel(
            referencia_id=uuid.uuid4(),
            filename="proyecto.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            content=content.read(),
        )

        await service.confirm_project_excel_import(
            referencia_id=uuid.uuid4(),
        )

        assert storage.document.storage_key.startswith("proyectos-formativos/")

    async def test_confirm_heals_stale_programa_id(self):
        stale_programa_id = uuid.uuid4()
        correct_programa_id = uuid.uuid4()
        programa_referencia_id = uuid.uuid4()
        proyecto_referencia_id = uuid.uuid4()

        proyecto_draft = MagicMock(
            id=uuid.uuid4(),
            payload_json={
                "meta": {
                    "referenciaId": str(proyecto_referencia_id),
                    "programaReferenciaId": str(programa_referencia_id),
                    "programaId": str(stale_programa_id),
                    "touchedSteps": ["fuente-proyecto"],
                    "lastInteractionAt": datetime.now(UTC).isoformat(),
                },
                "documental": {},
            },
            paso_actual="fuente-proyecto",
            estado_borrador=EstadoBloque.BORRADOR,
        )
        programa_draft = MagicMock(
            payload_json={
                "meta": {"referenciaId": str(programa_referencia_id)},
                "curricular": {"programa_formacion_id": str(correct_programa_id)},
            },
        )

        mock_execute_res = MagicMock()
        mock_execute_res.scalar_one_or_none.return_value = None
        mock_execute_res.scalars.return_value.unique.return_value.all.return_value = []

        session = MockSession()
        session.execute = MagicMock(return_value=mock_execute_res)

        draft_repository = MockDraftRepository(proyecto_draft, programa_draft)
        audit_repository = MockAuditRepository()
        storage_service = MockStorageService()
        project_repository = MockProjectRepository()

        service = ProyectoExcelImportService(
            session=session,
            draft_repository=draft_repository,
            audit_repository=audit_repository,
            project_repository=project_repository,
            storage_service=storage_service,
        )

        wb = create_valid_workbook()
        content = io.BytesIO()
        wb.save(content)
        content.seek(0)

        await service.preview_project_excel(
            referencia_id=proyecto_referencia_id,
            filename="proyecto.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            content=content.read(),
        )

        await service.confirm_project_excel_import(
            referencia_id=proyecto_referencia_id,
        )

        assert project_repository.created_programa_ids == [correct_programa_id]
        assert proyecto_draft.payload_json["meta"]["programaId"] == str(
            correct_programa_id
        )


_PLANEACION_HEADERS = [
    "proyecto_id",
    "fase_id",
    "fase_proyecto",
    "actividad_id",
    "actividad_proyecto",
    "tipo_resultado",
    "competencia_id",
    "codigo_competencia",
    "nombre_competencia",
    "rap_id",
    "rap_numero",
    "resultado_aprendizaje",
    "orden_fase",
    "orden_actividad",
    "orden_resultado",
    "pagina_origen",
    "observaciones",
]


def _append_planeacion_row(sheet, *, actividad_id, actividad, tipo, codigo_comp, comp_id, rap_id, rap_num, rap_desc, orden_resultado):
    sheet.append(
        [
            "PROJ-01",
            "F1",
            "Fase Analisis",
            actividad_id,
            actividad,
            tipo,
            comp_id,
            codigo_comp,
            f"Competencia {codigo_comp}",
            rap_id,
            rap_num,
            rap_desc,
            1,
            1,
            orden_resultado,
            "10",
            "",
        ]
    )


def create_multicompetencia_workbook():
    """One activity linked to two competencies; first competency has two RAPs."""

    wb = Workbook()
    proyecto_sheet = wb.active
    proyecto_sheet.title = "Proyecto"
    proyecto_sheet.append(
        [
            "proyecto_id",
            "nombre_proyecto",
            "codigo_proyecto_sofia",
            "codigo_programa",
            "nombre_programa",
            "fuente_archivo",
            "observaciones",
        ],
    )
    proyecto_sheet.append(
        [
            "PROJ-01",
            "Proyecto integrado",
            "PR-002",
            "prog-ref",
            "Programa integrado",
            "fuente.pdf",
            "",
        ]
    )

    planeacion_sheet = wb.create_sheet("Planeacion_Proyecto")
    planeacion_sheet.append(_PLANEACION_HEADERS)
    # Actividad unica con competencia tecnica (2 RAP) y transversal (1 RAP)
    _append_planeacion_row(
        planeacion_sheet,
        actividad_id="A1",
        actividad="Estructurar propuesta tecnica",
        tipo="ESPECIFICO",
        codigo_comp="220501094",
        comp_id="C1",
        rap_id="R10",
        rap_num="1",
        rap_desc="Definir especificaciones",
        orden_resultado=1,
    )
    _append_planeacion_row(
        planeacion_sheet,
        actividad_id="A1",
        actividad="Estructurar propuesta tecnica",
        tipo="ESPECIFICO",
        codigo_comp="220501094",
        comp_id="C1",
        rap_id="R11",
        rap_num="3",
        rap_desc="Validar condiciones",
        orden_resultado=3,
    )
    _append_planeacion_row(
        planeacion_sheet,
        actividad_id="A1",
        actividad="Estructurar propuesta tecnica",
        tipo="TRANSVERSAL",
        codigo_comp="240201524",
        comp_id="C2",
        rap_id="R20",
        rap_num="3",
        rap_desc="Relacionar procesos comunicativos",
        orden_resultado=1,
    )

    validacion_sheet = wb.create_sheet("Validacion_Proyecto")
    validacion_sheet.append(["tipo_validacion", "descripcion", "estado", "observaciones"])
    validacion_sheet.append(["CURRICULAR", "Validacion", "OK", ""])

    return wb


class TestMaterializacionCurricular:
    def _build_service(self, competencias):
        draft = MagicMock(
            id=uuid.uuid4(),
            payload_json={
                "meta": {
                    "referenciaId": str(uuid.uuid4()),
                    "programaId": "00000000-0000-0000-0000-000000000000",
                    "touchedSteps": ["fuente-proyecto"],
                    "lastInteractionAt": datetime.now(UTC).isoformat(),
                },
                "documental": {},
            },
            paso_actual="fuente-proyecto",
            estado_borrador=EstadoBloque.BORRADOR,
        )
        session = MockSession(competencias=competencias)
        service = ProyectoExcelImportService(
            session=session,
            draft_repository=MockDraftRepository(draft),
            audit_repository=MockAuditRepository(),
            project_repository=MockProjectRepository(),
            storage_service=MockStorageService(),
        )
        return service

    async def _preview_and_confirm(self, service, workbook):
        content = io.BytesIO()
        workbook.save(content)
        content.seek(0)
        await service.preview_project_excel(
            referencia_id=uuid.uuid4(),
            filename="proyecto.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            content=content.read(),
        )
        return await service.confirm_project_excel_import(referencia_id=uuid.uuid4())

    async def test_actividad_links_multiple_competencias(self):
        competencias = [
            _build_competencia(
                codigo_competencia="220501094",
                nombre="Estructurar propuesta",
                rap_codigos=["R10", "R11"],
            ),
            _build_competencia(
                codigo_competencia="240201524",
                nombre="Comunicacion",
                rap_codigos=["R20"],
            ),
        ]
        service = self._build_service(competencias)
        result = await self._preview_and_confirm(service, create_multicompetencia_workbook())

        repo = service._project_repository
        assert len(result.actividad_ids) == 1
        actividad_id = result.actividad_ids[0]
        actividad_asignaciones = [
            a for a in repo.created_asignaciones if a.actividad_proyecto_id == actividad_id
        ]
        competencias_ids = {a.competencia_id for a in actividad_asignaciones}
        assert len(competencias_ids) == 2
        assert result.pendientes_resumen.asignaciones_materializadas == 3
        assert result.pendientes_resumen.total == 0

    async def test_competencia_links_multiple_raps(self):
        competencias = [
            _build_competencia(
                codigo_competencia="220501094",
                nombre="Estructurar propuesta",
                rap_codigos=["R10", "R11"],
            ),
            _build_competencia(
                codigo_competencia="240201524",
                nombre="Comunicacion",
                rap_codigos=["R20"],
            ),
        ]
        service = self._build_service(competencias)
        await self._preview_and_confirm(service, create_multicompetencia_workbook())

        repo = service._project_repository
        comp_tecnica = competencias[0]
        raps_tecnica = [
            a for a in repo.created_asignaciones if a.competencia_id == comp_tecnica.id
        ]
        assert len(raps_tecnica) == 2
        assert {a.resultado_id for a in raps_tecnica} == {
            comp_tecnica.resultados[0].id,
            comp_tecnica.resultados[1].id,
        }

    async def test_tipo_resultado_is_preserved(self):
        competencias = [
            _build_competencia(
                codigo_competencia="220501094",
                nombre="Estructurar propuesta",
                rap_codigos=["R10", "R11"],
            ),
            _build_competencia(
                codigo_competencia="240201524",
                nombre="Comunicacion",
                rap_codigos=["R20"],
            ),
        ]
        service = self._build_service(competencias)
        await self._preview_and_confirm(service, create_multicompetencia_workbook())

        repo = service._project_repository
        tipos = sorted(a.tipo_resultado for a in repo.created_asignaciones)
        assert tipos == ["ESPECIFICO", "ESPECIFICO", "TRANSVERSAL"]

    async def test_missing_tipo_resultado_column_materializes_null(self):
        competencias = [
            _build_competencia(
                codigo_competencia="220501094",
                nombre="Estructurar propuesta",
                rap_codigos=["R10", "R11"],
            ),
            _build_competencia(
                codigo_competencia="240201524",
                nombre="Comunicacion",
                rap_codigos=["R20"],
            ),
        ]
        workbook = create_multicompetencia_workbook()
        workbook["Planeacion_Proyecto"].delete_cols(6)
        service = self._build_service(competencias)

        await self._preview_and_confirm(service, workbook)

        repo = service._project_repository
        assert len(repo.created_asignaciones) == 3
        assert all(a.tipo_resultado is None for a in repo.created_asignaciones)

    async def test_unresolved_competencia_is_reported(self):
        competencias = [
            _build_competencia(
                codigo_competencia="999999999",
                nombre="Otra competencia",
                rap_codigos=["RX"],
            ),
        ]
        service = self._build_service(competencias)
        result = await self._preview_and_confirm(service, create_multicompetencia_workbook())

        assert result.pendientes_resumen.asignaciones_sin_competencia == 3
        assert result.pendientes_resumen.asignaciones_materializadas == 0
        assert result.pendientes_resumen.total == 3

    async def test_unresolved_resultado_is_reported(self):
        competencias = [
            _build_competencia(
                codigo_competencia="220501094",
                nombre="Estructurar propuesta",
                rap_codigos=["R10"],
            ),
            _build_competencia(
                codigo_competencia="240201524",
                nombre="Comunicacion",
                rap_codigos=["R20"],
            ),
        ]
        service = self._build_service(competencias)
        result = await self._preview_and_confirm(service, create_multicompetencia_workbook())

        repo = service._project_repository
        # R11 no existe en la competencia tecnica: se reporta sin resultado
        assert result.pendientes_resumen.asignaciones_sin_resultado == 1
        sin_resultado = [a for a in repo.created_asignaciones if a.resultado_id is None]
        assert len(sin_resultado) == 1

    async def test_default_workbook_materializes_assignments(self, service):
        workbook = create_valid_workbook()
        result = await self._preview_and_confirm(service, workbook)

        repo = service._project_repository
        assert result.pendientes_resumen.asignaciones_materializadas == 3
        assert result.pendientes_resumen.total == 0
        assert len(repo.created_asignaciones) == 3


class TestTipoResultadoValidation:
    def _create_workbook_with_tipo_resultado(self, tipo_value: str | None) -> bytes:
        wb = openpyxl.Workbook()
        ws_proj = wb.active
        ws_proj.title = "Proyecto"
        ws_proj.append(
            [
                "proyecto_id",
                "nombre_proyecto",
                "codigo_proyecto_sofia",
                "codigo_programa",
                "nombre_programa",
                "fuente_archivo",
                "observaciones",
            ]
        )
        ws_proj.append(["PROJ-01", "Proyecto Test", "SOFIA-01", "228118", "Programa", "MATRIZ", ""])

        ws_plan = wb.create_sheet("Planeacion_Proyecto")
        ws_plan.append(_PLANEACION_HEADERS)
        _append_planeacion_row(
            ws_plan,
            actividad_id="A1",
            actividad="Actividad 1",
            tipo=tipo_value,
            codigo_comp="220501094",
            comp_id="C1",
            rap_id="R10",
            rap_num="1",
            rap_desc="Desc R10",
            orden_resultado=1,
        )

        ws_val = wb.create_sheet("Validacion_Proyecto")
        ws_val.append(["tipo_validacion", "descripcion", "estado", "observaciones"])
        ws_val.append(["CURRICULAR", "Validacion", "OK", ""])

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    async def test_tipo_resultado_especifico_normalized(self):
        content = self._create_workbook_with_tipo_resultado("  especifico  ")
        parsed = parse_canonical_workbook(content)
        assert len(parsed.errores) == 0
        assert len(parsed.planeacion) == 1
        assert parsed.planeacion[0].tipo_resultado == "ESPECIFICO"

    async def test_tipo_resultado_transversal_normalized(self):
        content = self._create_workbook_with_tipo_resultado("transversal")
        parsed = parse_canonical_workbook(content)
        assert len(parsed.errores) == 0
        assert len(parsed.planeacion) == 1
        assert parsed.planeacion[0].tipo_resultado == "TRANSVERSAL"

    async def test_tipo_resultado_basico_normalized(self):
        content = self._create_workbook_with_tipo_resultado("  basico  ")
        parsed = parse_canonical_workbook(content)
        assert len(parsed.errores) == 0
        assert len(parsed.planeacion) == 1
        assert parsed.planeacion[0].tipo_resultado == "BASICO"

    async def test_tipo_resultado_blank_is_optional(self):
        content = self._create_workbook_with_tipo_resultado(None)
        parsed = parse_canonical_workbook(content)
        assert len(parsed.errores) == 0
        assert len(parsed.planeacion) == 1
        assert parsed.planeacion[0].tipo_resultado is None

    async def test_tipo_resultado_column_is_optional(self):
        content = self._create_workbook_with_tipo_resultado("ESPECIFICO")
        workbook = openpyxl.load_workbook(io.BytesIO(content))
        workbook["Planeacion_Proyecto"].delete_cols(6)
        output = io.BytesIO()
        workbook.save(output)

        parsed = parse_canonical_workbook(output.getvalue())

        assert len(parsed.errores) == 0
        assert len(parsed.planeacion) == 1
        assert parsed.planeacion[0].tipo_resultado is None

    async def test_tipo_resultado_invalid_value_rejected(self):
        for invalid_val in ["TECNICO", "OTRO", "TRANSVERS", "INVALIDO"]:
            content = self._create_workbook_with_tipo_resultado(invalid_val)
            parsed = parse_canonical_workbook(content)
            assert len(parsed.planeacion) == 0
            assert len(parsed.errores) == 1
            issue = parsed.errores[0]
            assert issue.hoja == "Planeacion_Proyecto"
            assert issue.campo == "tipo_resultado"
            assert issue.fila == 2
            assert "Valor invalido para 'tipo_resultado'" in issue.mensaje

    async def test_tipo_resultado_metrics_counting(self):
        p1 = PlaneacionRow(
            proyecto_id="P1",
            fase_id="F1",
            fase_proyecto="Analisis",
            actividad_id="A1",
            actividad_proyecto="Act 1",
            tipo_resultado="ESPECIFICO",
            competencia_id="C1",
            codigo_competencia="220501094",
            nombre_competencia="Comp 1",
            rap_id="RAP-1",
            rap_numero="1",
            resultado_aprendizaje="Resultado 1",
            orden_fase=1,
            orden_actividad=1,
            orden_resultado=1,
            pagina_origen=None,
            observaciones=None,
            raw={},
        )
        p2 = PlaneacionRow(
            proyecto_id="P1",
            fase_id="F1",
            fase_proyecto="Analisis",
            actividad_id="A1",
            actividad_proyecto="Act 1",
            tipo_resultado="TRANSVERSAL",
            competencia_id="C2",
            codigo_competencia="240201524",
            nombre_competencia="Comp 2",
            rap_id="RAP-2",
            rap_numero="2",
            resultado_aprendizaje="Resultado 2",
            orden_fase=1,
            orden_actividad=1,
            orden_resultado=2,
            pagina_origen=None,
            observaciones=None,
            raw={},
        )
        count = count_resultados_especificos([p1, p2])
        assert count == 1
