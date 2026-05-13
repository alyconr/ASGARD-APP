"""Tests for TASK-08.5 canonical Excel preview and import."""

from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime
from typing import Any

import pytest
from openpyxl import Workbook

from src.application.dto.programa_documentos import StoredDocumentDTO
from src.application.services.programa_excel import (
    CANONICAL_SHEETS,
    InvalidProgramaExcelUploadError,
    ProgramaExcelImportService,
    ProgramaExcelValidationError,
)
from src.domain.drafts.types import TipoBloqueBorrador
from src.domain.shared.enums import (
    EstadoBloque,
    EstadoCampo,
    EstadoConciliacionPendiente,
    TipoConocimiento,
)
from src.infrastructure.db.models.curriculum import (
    Competencia,
    Conocimiento,
    CriterioEvaluacion,
    ElementoCurricularPendiente,
    ProgramaFormacion,
    ResultadoAprendizaje,
)
from src.infrastructure.db.models.drafts import BorradorSesion


class FakeSession:
    """Minimal async session double used by the Excel service."""

    def __init__(self) -> None:
        """Track commit calls."""
        self.commits = 0

    async def commit(self) -> None:
        """Record a commit invocation."""
        self.commits += 1

    async def refresh(self, instance: object) -> None:
        """Populate ids and timestamps expected after database flush."""
        now = datetime.now(UTC)
        if hasattr(instance, "id") and getattr(instance, "id", None) is None:
            setattr(instance, "id", uuid.uuid4())
        if hasattr(instance, "fecha_creacion"):
            setattr(instance, "fecha_creacion", now)
        if hasattr(instance, "fecha_actualizacion"):
            setattr(instance, "fecha_actualizacion", now)
        if isinstance(instance, BorradorSesion):
            instance.ultima_edicion = now


class FakeDraftRepository:
    """In-memory draft repository keyed by stable reference."""

    def __init__(self, draft: BorradorSesion) -> None:
        """Store a single draft."""
        self.draft = draft

    async def get_by_block_reference(
        self,
        tipo_bloque: TipoBloqueBorrador,
        referencia_id: uuid.UUID,
    ) -> BorradorSesion | None:
        """Return the draft when the logical key matches."""
        if (
            tipo_bloque is TipoBloqueBorrador.PROGRAMA
            and self.draft.referencia_id == referencia_id
        ):
            return self.draft
        return None

    async def save(self, draft: BorradorSesion) -> BorradorSesion:
        """Persist draft changes."""
        self.draft = draft
        return draft


class FakeAuditRepository:
    """Collect audit events."""

    def __init__(self) -> None:
        """Initialize empty audit storage."""
        self.events: list[dict[str, object]] = []

    async def add_event(
        self,
        entidad: str,
        entidad_id: uuid.UUID,
        accion: str,
        detalle: dict[str, object] | None = None,
    ) -> object:
        """Append an audit event."""
        event: dict[str, object] = {
            "entidad": entidad,
            "entidad_id": entidad_id,
            "accion": accion,
            "detalle": detalle or {},
        }
        self.events.append(event)
        return event


class FakeStorageService:
    """In-memory object storage for Excel files."""

    def __init__(self) -> None:
        """Initialize empty object storage."""
        self.objects: dict[str, bytes] = {}

    async def save_excel(
        self,
        *,
        key: str,
        content: bytes,
        content_type: str,
        original_filename: str,
    ) -> StoredDocumentDTO:
        """Store bytes and return deterministic metadata."""
        self.objects[key] = content
        return StoredDocumentDTO(
            original_filename=original_filename,
            storage_key=key,
            size_bytes=len(content),
            content_type=content_type,
            checksum_sha256="xlsx-checksum",
            etag="etag-xlsx",
        )

    async def read_excel(self, *, key: str) -> bytes:
        """Read stored bytes."""
        try:
            return self.objects[key]
        except KeyError as error:
            raise FileNotFoundError(key) from error


class FakeProgramaExcelRepository:
    """In-memory repository for all curriculum rows."""

    def __init__(self) -> None:
        """Initialize empty storage."""
        self.programas: dict[uuid.UUID, ProgramaFormacion] = {}
        self.competencias: dict[uuid.UUID, Competencia] = {}
        self.resultados: dict[uuid.UUID, ResultadoAprendizaje] = {}
        self.conocimientos: dict[uuid.UUID, Conocimiento] = {}
        self.criterios: dict[uuid.UUID, CriterioEvaluacion] = {}
        self.pendientes: dict[uuid.UUID, ElementoCurricularPendiente] = {}

    async def get_programa(self, programa_id: uuid.UUID) -> ProgramaFormacion | None:
        """Return a program by id."""
        return self.programas.get(programa_id)

    async def get_programa_by_code_version(
        self,
        codigo_programa: str,
        version_programa: str | None,
    ) -> ProgramaFormacion | None:
        """Return matching program."""
        for programa in self.programas.values():
            if (
                programa.codigo_programa.lower() == codigo_programa.lower()
                and programa.version_programa == version_programa
            ):
                return programa
        return None

    async def add_programa(
        self,
        *,
        codigo_programa: str,
        nombre_programa: str,
        version_programa: str | None,
    ) -> ProgramaFormacion:
        """Create a program."""
        now = datetime.now(UTC)
        programa = ProgramaFormacion(
            codigo_programa=codigo_programa,
            nombre_programa=nombre_programa,
            version_programa=version_programa,
        )
        programa.id = uuid.uuid4()
        programa.fecha_creacion = now
        programa.fecha_actualizacion = now
        self.programas[programa.id] = programa
        return programa

    async def list_competencias(self, programa_id: uuid.UUID) -> list[Competencia]:
        """List competences for a program."""
        competencias = [
            competencia
            for competencia in self.competencias.values()
            if competencia.programa_id == programa_id
        ]
        for competencia in competencias:
            competencia.resultados = [
                item
                for item in self.resultados.values()
                if item.competencia_id == competencia.id
            ]
            competencia.conocimientos = [
                item
                for item in self.conocimientos.values()
                if item.competencia_id == competencia.id
            ]
            competencia.criterios = [
                item
                for item in self.criterios.values()
                if item.competencia_id == competencia.id
            ]
        return competencias

    async def competencia_code_exists(
        self,
        *,
        programa_id: uuid.UUID,
        codigo_competencia: str,
    ) -> bool:
        """Check duplicate competence code."""
        return any(
            competencia.programa_id == programa_id
            and competencia.codigo_competencia.lower() == codigo_competencia.lower()
            for competencia in self.competencias.values()
        )

    async def resultado_exists(
        self,
        *,
        competencia_id: uuid.UUID,
        descripcion: str,
    ) -> bool:
        """Check duplicate result."""
        return any(
            item.competencia_id == competencia_id
            and item.descripcion.lower() == descripcion.lower()
            for item in self.resultados.values()
        )

    async def conocimiento_exists(
        self,
        *,
        competencia_id: uuid.UUID,
        tipo: TipoConocimiento,
        descripcion: str,
        resultado_id: uuid.UUID | None = None,
    ) -> bool:
        """Check duplicate knowledge."""
        return any(
            item.competencia_id == competencia_id
            and item.tipo == tipo
            and item.descripcion.lower() == descripcion.lower()
            and item.resultado_id == resultado_id
            for item in self.conocimientos.values()
        )

    async def criterio_exists(
        self,
        *,
        competencia_id: uuid.UUID,
        descripcion: str,
        resultado_id: uuid.UUID | None = None,
    ) -> bool:
        """Check duplicate criterion."""
        return any(
            item.competencia_id == competencia_id
            and item.descripcion.lower() == descripcion.lower()
            and item.resultado_id == resultado_id
            for item in self.criterios.values()
        )

    async def add_competencia(
        self,
        *,
        programa_id: uuid.UUID,
        codigo_competencia: str,
        nombre_competencia: str,
        orden: int | None,
    ) -> Competencia:
        """Create a competence."""
        now = datetime.now(UTC)
        competencia = Competencia(
            programa_id=programa_id,
            codigo_competencia=codigo_competencia,
            nombre_competencia=nombre_competencia,
            orden=orden,
            estado=EstadoBloque.BORRADOR,
            origen_campo=EstadoCampo.VALIDADO,
        )
        competencia.id = uuid.uuid4()
        competencia.fecha_creacion = now
        competencia.fecha_actualizacion = now
        self.competencias[competencia.id] = competencia
        return competencia

    async def add_resultado(
        self,
        *,
        competencia_id: uuid.UUID,
        codigo_resultado: str | None,
        descripcion: str,
        orden: int | None,
    ) -> ResultadoAprendizaje:
        """Create a result."""
        resultado = ResultadoAprendizaje(
            competencia_id=competencia_id,
            codigo_resultado=codigo_resultado,
            descripcion=descripcion,
            orden=orden,
            estado=EstadoCampo.VALIDADO,
        )
        resultado.id = uuid.uuid4()
        self.resultados[resultado.id] = resultado
        return resultado

    async def add_conocimiento(
        self,
        *,
        competencia_id: uuid.UUID,
        tipo: TipoConocimiento,
        descripcion: str,
        orden: int | None,
        resultado_id: uuid.UUID | None = None,
    ) -> Conocimiento:
        """Create a knowledge row."""
        conocimiento = Conocimiento(
            competencia_id=competencia_id,
            resultado_id=resultado_id,
            tipo=tipo,
            descripcion=descripcion,
            orden=orden,
            estado=EstadoCampo.VALIDADO,
        )
        conocimiento.id = uuid.uuid4()
        self.conocimientos[conocimiento.id] = conocimiento
        return conocimiento

    async def add_criterio(
        self,
        *,
        competencia_id: uuid.UUID,
        descripcion: str,
        orden: int | None,
        resultado_id: uuid.UUID | None = None,
    ) -> CriterioEvaluacion:
        """Create a criterion."""
        criterio = CriterioEvaluacion(
            competencia_id=competencia_id,
            resultado_id=resultado_id,
            descripcion=descripcion,
            orden=orden,
            estado=EstadoCampo.VALIDADO,
        )
        criterio.id = uuid.uuid4()
        self.criterios[criterio.id] = criterio
        return criterio

    async def clear_pendientes(self, *, referencia_id: uuid.UUID) -> None:
        """Remove pending rows for a reference."""
        self.pendientes = {
            key: value
            for key, value in self.pendientes.items()
            if value.referencia_id != referencia_id
        }

    async def add_pendiente(
        self,
        *,
        referencia_id: uuid.UUID,
        programa_id: uuid.UUID,
        pendiente: Any,
        orden: int | None,
        raw_excel: dict[str, object] | None,
    ) -> ElementoCurricularPendiente:
        """Create a pending row."""
        now = datetime.now(UTC)
        item = ElementoCurricularPendiente(
            referencia_id=referencia_id,
            programa_id=programa_id,
            tipo_elemento=pendiente.tipo_elemento,
            tipo_conocimiento=pendiente.tipo_conocimiento,
            descripcion=pendiente.descripcion,
            competencia_id_origen_excel=pendiente.competencia_id_origen_excel,
            rap_id_origen_excel=pendiente.rap_id_origen_excel,
            motivo=pendiente.motivo,
            estado=EstadoConciliacionPendiente.PENDIENTE,
            orden=orden,
            raw_excel=raw_excel,
        )
        item.id = uuid.uuid4()
        item.fecha_creacion = now
        item.fecha_actualizacion = now
        self.pendientes[item.id] = item
        return item


def build_draft(referencia_id: uuid.UUID) -> BorradorSesion:
    """Create a program draft ORM object."""
    draft = BorradorSesion(
        tipo_bloque=TipoBloqueBorrador.PROGRAMA.value,
        referencia_id=referencia_id,
        paso_actual="origen-documental",
        payload_json={
            "meta": {
                "referenciaId": str(referencia_id),
                "entryMode": "EXCEL",
                "touchedSteps": ["datos-programa"],
            },
            "programa": {
                "codigo_programa": "",
                "nombre_programa": "",
                "version_programa": "",
            },
            "documental": {"programa_pdf": None, "programa_excel": None},
            "curricular": {"programa_formacion_id": None, "competencias": []},
        },
        estado_borrador=EstadoBloque.BORRADOR,
    )
    draft.id = uuid.uuid4()
    draft.ultima_edicion = datetime.now(UTC)
    return draft


def build_service() -> tuple[
    ProgramaExcelImportService,
    FakeSession,
    FakeDraftRepository,
    FakeProgramaExcelRepository,
    FakeStorageService,
]:
    """Build the service with fake dependencies."""
    referencia_id = uuid.uuid4()
    session = FakeSession()
    drafts = FakeDraftRepository(build_draft(referencia_id))
    repo = FakeProgramaExcelRepository()
    storage = FakeStorageService()
    service = ProgramaExcelImportService(
        session=session,
        draft_repository=drafts,
        audit_repository=FakeAuditRepository(),
        curriculum_repository=repo,
        storage_service=storage,
    )
    return service, session, drafts, repo, storage


def build_workbook_bytes(overrides: dict[str, list[list[Any]]] | None = None) -> bytes:
    """Create a canonical workbook in memory."""
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    rows = {
        "Programa": [
            [
                "228118",
                "Analisis y desarrollo de software",
                "1",
                "2026",
                "Tecnologo",
                "Tecnologo ADSO",
                864,
                432,
                1296,
                "programa.pdf",
                "",
            ],
        ],
        "Competencias": [
            ["COMP-1", "220501046", "Desarrollar software", 120, 1, 10, ""],
        ],
        "Resultados": [
            ["COMP-1", "RAP-1", "1", "Construye componentes", 1, 11, ""],
        ],
        "Conocimientos": [
            ["COMP-1", "RAP-1", "SABER", "1", "Arquitectura", 1, 12, ""],
            ["COMP-1", "", "PROCESO", "1", "Codificar solucion", 2, 13, ""],
        ],
        "Criterios": [
            ["COMP-1", "RAP-1", "1", "Verifica componentes", 1, 14, ""],
        ],
    }
    if overrides:
        rows.update(overrides)

    for sheet_name, headers in CANONICAL_SHEETS.items():
        sheet = workbook.create_sheet(title=sheet_name)
        sheet.append(headers)
        for row in rows.get(sheet_name, []):
            sheet.append(row)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


@pytest.mark.anyio
async def test_preview_rejects_non_excel_file() -> None:
    """Only .xlsx files are accepted."""
    service, _, drafts, _, _ = build_service()

    with pytest.raises(InvalidProgramaExcelUploadError):
        await service.preview_program_excel(
            referencia_id=drafts.draft.referencia_id,
            filename="programa.csv",
            content_type="text/csv",
            content=b"not excel",
        )


@pytest.mark.anyio
async def test_preview_reports_invalid_workbook_headers() -> None:
    """Invalid canonical structure should return validation errors."""
    service, _, drafts, _, storage = build_service()
    content = build_workbook_bytes(overrides={"Competencias": []})
    # Corrupt the first expected header.
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, headers in CANONICAL_SHEETS.items():
        sheet = workbook.create_sheet(title=sheet_name)
        sheet.append(["bad"] + headers[1:] if sheet_name == "Competencias" else headers)
    output = io.BytesIO()
    workbook.save(output)

    result = await service.preview_program_excel(
        referencia_id=drafts.draft.referencia_id,
        filename="programa.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=output.getvalue(),
    )

    assert result.valid is False
    assert result.errores[0].hoja == "Competencias"
    assert storage.objects == {}
    assert content


@pytest.mark.anyio
async def test_preview_valid_canonical_workbook_updates_same_draft() -> None:
    """A valid workbook should be stored and previewed without relational writes."""
    service, session, drafts, repo, storage = build_service()
    referencia_id = drafts.draft.referencia_id

    result = await service.preview_program_excel(
        referencia_id=referencia_id,
        filename="programa.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=build_workbook_bytes(),
    )

    assert result.referencia_id == referencia_id
    assert result.valid is True
    assert result.resumen.competencias == 1
    assert result.resumen.resultados == 1
    assert result.resumen.conocimientos == 2
    assert result.pendientes_resumen.conocimientos == 0
    assert result.competencias[0].resultados_detalle[0].rap_id == "RAP-1"
    assert len(result.competencias[0].conocimientos_detalle) == 2
    assert len(result.competencias[0].criterios_detalle) == 1
    assert repo.programas == {}
    assert len(storage.objects) == 1
    assert drafts.draft.referencia_id == referencia_id
    assert (
        drafts.draft.payload_json["documental"]["programa_excel"]["preview"]["valid"]
        is True
    )
    assert session.commits == 1


@pytest.mark.anyio
async def test_preview_rejects_broken_cross_references() -> None:
    """Children cannot point to absent competencia_id values."""
    service, _, drafts, _, _ = build_service()

    result = await service.preview_program_excel(
        referencia_id=drafts.draft.referencia_id,
        filename="programa.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=build_workbook_bytes(
            overrides={
                "Resultados": [
                    ["NO-EXISTE", "RAP-1", "1", "Construye componentes", 1, 11, ""],
                ],
            },
        ),
    )

    assert result.valid is False
    assert any("competencia_id no existe" in issue.mensaje for issue in result.errores)


@pytest.mark.anyio
async def test_preview_keeps_conocimiento_without_competencia_as_pending() -> None:
    """Knowledge rows without competencia_id do not invalidate the workbook."""
    service, _, drafts, _, _ = build_service()

    result = await service.preview_program_excel(
        referencia_id=drafts.draft.referencia_id,
        filename="programa.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=build_workbook_bytes(
            overrides={
                "Conocimientos": [
                    ["", "RAP-1", "SABER", "1", "Arquitectura pendiente", 1, 12, ""],
                ],
            },
        ),
    )

    assert result.valid is True
    assert result.errores == []
    assert result.resumen.conocimientos == 0
    assert result.pendientes_resumen.conocimientos == 1
    assert result.pendientes[0].motivo.value == "COMPETENCIA_NO_IDENTIFICADA"


@pytest.mark.anyio
async def test_preview_keeps_criterio_without_competencia_as_pending() -> None:
    """Criterion rows without competencia_id remain true reconciliation pending."""
    service, _, drafts, _, _ = build_service()

    result = await service.preview_program_excel(
        referencia_id=drafts.draft.referencia_id,
        filename="programa.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=build_workbook_bytes(
            overrides={
                "Criterios": [
                    ["", "", "1", "Criterio sin competencia", 1, 14, ""],
                ],
            },
        ),
    )

    assert result.valid is True
    assert result.errores == []
    assert result.resumen.criterios == 0
    assert result.pendientes_resumen.criterios == 1
    assert result.pendientes[0].motivo.value == "COMPETENCIA_NO_IDENTIFICADA"


@pytest.mark.anyio
async def test_preview_imports_criterio_with_competencia_without_rap() -> None:
    """Criteria rows with competence and without rap_id stay under competence."""
    service, _, drafts, _, _ = build_service()

    result = await service.preview_program_excel(
        referencia_id=drafts.draft.referencia_id,
        filename="programa.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=build_workbook_bytes(
            overrides={
                "Criterios": [
                    ["COMP-1", "", "1", "Criterio pendiente", 1, 14, ""],
                ],
            },
        ),
    )

    assert result.valid is True
    assert result.errores == []
    assert result.resumen.criterios == 1
    assert result.pendientes_resumen.criterios == 0
    assert result.competencias[0].criterios_detalle[0].descripcion == (
        "Criterio pendiente"
    )


@pytest.mark.anyio
async def test_preview_allows_practical_stage_competence_without_children() -> None:
    """A practical-stage competence may exist without curricular children."""
    service, _, drafts, _, _ = build_service()

    result = await service.preview_program_excel(
        referencia_id=drafts.draft.referencia_id,
        filename="programa.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=build_workbook_bytes(
            overrides={
                "Competencias": [
                    ["COMP-1", "220501046", "Desarrollar software", 120, 1, 10, ""],
                    [
                        "ETAPA-PRACTICA",
                        "999999999",
                        "Aplicar etapa practica",
                        864,
                        2,
                        20,
                        "",
                    ],
                ],
            },
        ),
    )

    assert result.valid is True
    practical = next(
        item for item in result.competencias if item.competencia_id == "ETAPA-PRACTICA"
    )
    assert practical.resultados == 0
    assert practical.conocimientos == 0
    assert practical.criterios == 0


@pytest.mark.anyio
async def test_confirm_import_materializes_full_curriculum_without_new_reference() -> (
    None
):
    """Confirmation should persist program, children and keep referencia_id stable."""
    service, session, drafts, repo, _ = build_service()
    referencia_id = drafts.draft.referencia_id
    await service.preview_program_excel(
        referencia_id=referencia_id,
        filename="programa.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=build_workbook_bytes(),
    )

    result = await service.confirm_program_excel_import(referencia_id=referencia_id)

    assert result.referencia_id == referencia_id
    assert len(repo.programas) == 1
    assert len(repo.competencias) == 1
    assert len(repo.resultados) == 1
    assert len(repo.conocimientos) == 2
    assert len(repo.criterios) == 1
    assert len(repo.pendientes) == 0
    assert drafts.draft.referencia_id == referencia_id
    assert drafts.draft.payload_json["curricular"]["programa_formacion_id"] == str(
        result.programa_id,
    )
    assert (
        drafts.draft.payload_json["documental"]["programa_excel"]["confirmacion"][
            "estado"
        ]
        == "IMPORTADO"
    )
    assert session.commits == 2


@pytest.mark.anyio
async def test_confirm_import_rejects_duplicate_confirmation() -> None:
    """The same preview cannot be confirmed twice and create duplicates."""
    service, _, drafts, _, _ = build_service()
    referencia_id = drafts.draft.referencia_id
    await service.preview_program_excel(
        referencia_id=referencia_id,
        filename="programa.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=build_workbook_bytes(),
    )
    await service.confirm_program_excel_import(referencia_id=referencia_id)

    with pytest.raises(ProgramaExcelValidationError):
        await service.confirm_program_excel_import(referencia_id=referencia_id)


@pytest.mark.anyio
async def test_preview_rejects_duplicate_resultado_rap_id() -> None:
    """RAP identifiers are unique inside a competence in the canonical workbook."""
    service, _, drafts, _, _ = build_service()

    result = await service.preview_program_excel(
        referencia_id=drafts.draft.referencia_id,
        filename="programa.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=build_workbook_bytes(
            overrides={
                "Resultados": [
                    ["COMP-1", "RAP-1", "1", "Construye componentes", 1, 11, ""],
                    ["COMP-1", "RAP-1", "2", "Integra componentes", 2, 12, ""],
                ],
            },
        ),
    )

    assert result.valid is False
    assert any(
        "Resultado duplicado por rap_id" in issue.mensaje for issue in result.errores
    )


@pytest.mark.anyio
async def test_preview_moves_duplicate_conocimiento_in_competencia_to_pending() -> None:
    """The same knowledge description is ambiguous in the same competence."""
    service, _, drafts, _, _ = build_service()

    result = await service.preview_program_excel(
        referencia_id=drafts.draft.referencia_id,
        filename="programa.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=build_workbook_bytes(
            overrides={
                "Resultados": [
                    ["COMP-1", "RAP-1", "1", "Construye componentes", 1, 11, ""],
                    ["COMP-1", "RAP-2", "2", "Integra componentes", 2, 12, ""],
                ],
                "Conocimientos": [
                    ["COMP-1", "RAP-1", "SABER", "1", "Arquitectura", 1, 12, ""],
                    ["COMP-1", "RAP-2", "SABER", "1", "Arquitectura", 1, 13, ""],
                ],
            },
        ),
    )

    assert result.valid is True
    assert result.resumen.conocimientos == 1
    assert result.pendientes_resumen.conocimientos == 1
    assert result.pendientes[0].motivo.value == "ASOCIACION_AMBIGUA"


@pytest.mark.anyio
async def test_preview_moves_duplicate_conocimiento_same_rap_to_pending() -> None:
    """The same knowledge description in the same RAP waits for reconciliation."""
    service, _, drafts, _, _ = build_service()

    result = await service.preview_program_excel(
        referencia_id=drafts.draft.referencia_id,
        filename="programa.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=build_workbook_bytes(
            overrides={
                "Conocimientos": [
                    ["COMP-1", "RAP-1", "SABER", "1", "Arquitectura", 1, 12, ""],
                    ["COMP-1", "RAP-1", "SABER", "2", "Arquitectura", 2, 13, ""],
                ],
            },
        ),
    )

    assert result.valid is True
    assert result.pendientes_resumen.conocimientos == 1
    assert result.pendientes[0].motivo.value == "ASOCIACION_AMBIGUA"


@pytest.mark.anyio
async def test_preview_moves_duplicate_criterio_in_competencia_to_pending() -> None:
    """The same criterion description is ambiguous in the same competence."""
    service, _, drafts, _, _ = build_service()

    result = await service.preview_program_excel(
        referencia_id=drafts.draft.referencia_id,
        filename="programa.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=build_workbook_bytes(
            overrides={
                "Resultados": [
                    ["COMP-1", "RAP-1", "1", "Construye componentes", 1, 11, ""],
                    ["COMP-1", "RAP-2", "2", "Integra componentes", 2, 12, ""],
                ],
                "Criterios": [
                    ["COMP-1", "RAP-1", "1", "Verifica componentes", 1, 14, ""],
                    ["COMP-1", "RAP-2", "1", "Verifica componentes", 1, 15, ""],
                ],
            },
        ),
    )

    assert result.valid is True
    assert result.resumen.criterios == 1
    assert result.pendientes_resumen.criterios == 1
    assert any(item.motivo.value == "ASOCIACION_AMBIGUA" for item in result.pendientes)


@pytest.mark.anyio
async def test_preview_moves_duplicate_criterio_same_rap_to_pending() -> None:
    """The same criterion description in the same RAP waits for reconciliation."""
    service, _, drafts, _, _ = build_service()

    result = await service.preview_program_excel(
        referencia_id=drafts.draft.referencia_id,
        filename="programa.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=build_workbook_bytes(
            overrides={
                "Criterios": [
                    ["COMP-1", "RAP-1", "1", "Verifica componentes", 1, 14, ""],
                    ["COMP-1", "RAP-1", "2", "Verifica componentes", 2, 15, ""],
                ],
            },
        ),
    )

    assert result.valid is True
    assert result.pendientes_resumen.criterios == 1
    assert any(item.motivo.value == "ASOCIACION_AMBIGUA" for item in result.pendientes)


@pytest.mark.anyio
async def test_preview_imports_items_without_rap_when_competencia_is_clear() -> None:
    """Items without rap_id should stay importable under their competence."""
    service, _, drafts, _, _ = build_service()

    result = await service.preview_program_excel(
        referencia_id=drafts.draft.referencia_id,
        filename="programa.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=build_workbook_bytes(
            overrides={
                "Conocimientos": [
                    ["COMP-1", "", "PROCESO", "1", "Codificar solucion", 2, 13, ""],
                    ["COMP-1", "", "PROCESO", "2", "Codificar solucion", 3, 14, ""],
                ],
                "Criterios": [
                    ["COMP-1", "", "1", "Prueba global", 1, 14, ""],
                    ["COMP-1", "", "2", "Prueba global", 2, 15, ""],
                ],
            },
        ),
    )

    assert result.valid is True
    assert result.errores == []
    assert result.resumen.conocimientos == 1
    assert result.resumen.criterios == 1
    assert result.pendientes_resumen.conocimientos == 1
    assert result.pendientes_resumen.criterios == 1


@pytest.mark.anyio
async def test_confirm_import_persists_resultado_id_correctly() -> None:
    """Confirmation should correctly resolve rap_id to the database resultado_id."""
    service, _, drafts, repo, _ = build_service()
    referencia_id = drafts.draft.referencia_id

    await service.preview_program_excel(
        referencia_id=referencia_id,
        filename="programa.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=build_workbook_bytes(),
    )

    await service.confirm_program_excel_import(referencia_id=referencia_id)

    # Find the result
    assert len(repo.resultados) == 1
    resultado = list(repo.resultados.values())[0]
    assert resultado.codigo_resultado == "RAP-1"

    # Verify conocimientos
    assert len(repo.conocimientos) == 2
    conocimientos = list(repo.conocimientos.values())

    saber = next(c for c in conocimientos if c.tipo == TipoConocimiento.SABER)

    # Saber knowledge should have the resultado_id since it was linked to RAP-1
    assert saber.resultado_id == resultado.id

    proceso = next(c for c in conocimientos if c.tipo == TipoConocimiento.PROCESO)

    # Proceso knowledge has no rap_id in the source and stays attached to competence.
    assert proceso.resultado_id is None
    assert len(repo.pendientes) == 0

    # Verify criterios
    assert len(repo.criterios) == 1
    criterio = list(repo.criterios.values())[0]

    # Criterion should have the resultado_id since it was linked to RAP-1
    assert criterio.resultado_id == resultado.id
